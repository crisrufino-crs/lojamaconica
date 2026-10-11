"""Importação do cadastro de Irmãos a partir de uma planilha do Excel (.xlsx).

Funciona em duas etapas: primeiro a planilha é só LIDA e conferida (nada é
gravado) e a tela mostra o que vai acontecer; a gravação só ocorre quando a
pessoa confirma.
"""
import io
import re
import unicodedata
from datetime import date, datetime

from django.db import transaction

from . import servicos
from .models import UFS, CampoExtra, Dependente, Grau, HistoricoGrau, Irmao

MAX_LINHAS = 5000
MAX_BYTES = 5 * 1024 * 1024


def chave(texto):
    """'Data de Nascimento ' -> 'datadenascimento' (sem acento, espaço ou pontuação)."""
    sem_acento = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", sem_acento.lower())


# (campo, título da coluna no modelo, tipo, outros títulos aceitos)
COLUNAS = [
    ("nome", "Nome completo", "texto", ["nome", "nome do irmao", "irmao", "nome completo do irmao"]),
    ("cim", "CIM", "codigo", ["cadastro", "cim cadastro", "numero cim", "n cim", "no cim", "num cim", "cim no", "matricula", "registro"]),
    ("situacao", "Situação", "situacao", ["status", "situacao cadastral"]),
    ("grau", "Grau", "grau", ["grau atual", "grau simbolico"]),
    ("cargo", "Cargo", "texto", ["cargo na loja", "funcao"]),
    ("data_nascimento", "Data de nascimento", "data", ["nascimento", "dt nascimento", "data nasc", "nasc", "aniversario"]),
    ("cpf", "CPF", "cpf", []),
    ("rg", "RG", "codigo", ["identidade", "rg identidade"]),
    ("email", "E-mail", "email", ["email", "e mail", "correio eletronico"]),
    ("celular", "Celular", "fone", ["telefone celular", "whatsapp", "cel", "fone celular"]),
    ("telefone", "Telefone", "fone", ["telefone fixo", "fone", "telefone residencial", "tel"]),
    ("endereco", "Endereço", "texto", ["endereco residencial", "logradouro", "rua"]),
    ("bairro", "Bairro", "texto", []),
    ("cidade", "Cidade", "texto", ["municipio"]),
    ("uf", "UF", "uf", ["estado"]),
    ("cep", "CEP", "cep", []),
    ("naturalidade", "Naturalidade", "texto", ["natural de", "cidade de nascimento"]),
    ("profissao", "Profissão", "texto", ["ocupacao"]),
    ("estado_civil", "Estado civil", "estado_civil", []),
    ("nome_pai", "Nome do pai", "texto", ["pai", "filiacao pai"]),
    ("nome_mae", "Nome da mãe", "texto", ["mae", "filiacao mae"]),
    ("data_iniciacao", "Data de iniciação", "data", ["iniciacao", "dt iniciacao", "iniciado em"]),
    ("data_elevacao", "Data de elevação", "data", ["elevacao", "dt elevacao", "elevado em"]),
    ("data_exaltacao", "Data de exaltação", "data", ["exaltacao", "dt exaltacao", "exaltado em"]),
    ("data_instalacao", "Data de instalação", "data", ["instalacao", "dt instalacao", "instalado em"]),
    ("data_filiacao", "Data de filiação à Loja", "data", ["filiacao", "data de filiacao", "dt filiacao", "filiado em"]),
    ("conjuge", "Cônjuge", "texto", ["esposa", "nome da esposa", "nome do conjuge", "cunhada"]),
    ("conjuge_nascimento", "Nascimento do cônjuge", "data", ["nascimento da esposa", "data de nascimento do conjuge",
                                                              "data de nascimento da esposa", "aniversario da esposa"]),
    ("observacoes", "Observações", "texto", ["obs", "observacao", "anotacoes"]),
]
TITULOS = {c[0]: c[1] for c in COLUNAS}
TIPOS = {c[0]: c[2] for c in COLUNAS}
POR_CHAVE = {}
for _campo, _titulo, _tipo, _outros in COLUNAS:
    for _t in [_titulo, _campo, *_outros]:
        POR_CHAVE.setdefault(chave(_t), _campo)

# Datas de grau -> posição do grau na Loja (1 = primeiro grau)
GRAU_POR_DATA = {"data_iniciacao": (1, "Iniciação"), "data_elevacao": (2, "Elevação"),
                 "data_exaltacao": (3, "Exaltação"), "data_instalacao": (4, "Instalação")}
TAMANHOS = {f.name: f.max_length for f in Irmao._meta.fields if getattr(f, "max_length", None)}
SITUACOES = {chave(r): v for v, r in Irmao.SITUACOES} | {chave(v): v for v, _ in Irmao.SITUACOES} | {
    "regular": "ativo", "quiteplacet": "quite_placet", "quite": "quite_placet", "inativo": "desligado",
    "adormecido": "desligado", "excluido": "desligado", "falecidos": "falecido", "orienteeterno": "falecido"}
ESTADOS_CIVIS = {chave(r): v for v, r in Irmao.ESTADOS_CIVIS} | {
    "casada": "casado", "solteira": "solteiro", "divorciada": "divorciado", "viuva": "viuvo", "separado": "divorciado",
    "amasiado": "uniao", "uniaoestavel": "uniao"}
NOMES_UF = {"acre": "AC", "alagoas": "AL", "amapa": "AP", "amazonas": "AM", "bahia": "BA", "ceara": "CE",
            "distritofederal": "DF", "espiritosanto": "ES", "goias": "GO", "maranhao": "MA", "matogrosso": "MT",
            "matogrossodosul": "MS", "minasgerais": "MG", "para": "PA", "paraiba": "PB", "parana": "PR",
            "pernambuco": "PE", "piaui": "PI", "riodejaneiro": "RJ", "riograndedonorte": "RN", "riograndedosul": "RS",
            "rondonia": "RO", "roraima": "RR", "santacatarina": "SC", "saopaulo": "SP", "sergipe": "SE", "tocantins": "TO"}
SIGLAS_UF = {u for u, _ in UFS}


class ErroDePlanilha(Exception):
    pass


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)  # o Excel guarda 123 como 123.0
    return re.sub(r"\s+", " ", str(valor)).strip()


def _so_digitos(valor):
    return re.sub(r"\D", "", _texto(valor))


def _data(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _texto(valor)
    if not texto:
        return None
    for formato in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%Y-%m-%d %H:%M:%S"):
        try:
            convertida = datetime.strptime(texto, formato).date()
        except ValueError:
            continue
        if formato == "%d/%m/%y" and convertida > date.today():
            convertida = convertida.replace(year=convertida.year - 100)  # "15/03/58" é 1958, não 2058
        return convertida
    raise ValueError(f"data “{texto}” não reconhecida (use dia/mês/ano, ex.: 25/03/1975)")


def converter(campo, tipo, valor, graus):
    """Transforma o conteúdo de uma célula no valor que vai para o cadastro. Levanta ValueError com o motivo."""
    texto = _texto(valor)
    if tipo == "data":
        convertida = _data(valor)
        if convertida and not (date(1850, 1, 1) <= convertida <= date.today()):
            raise ValueError(f"data {convertida:%d/%m/%Y} fora do esperado")
        return convertida
    if not texto:
        return ""
    if tipo == "cpf":
        digitos = _so_digitos(valor)
        if len(digitos) > 11 or not digitos:
            raise ValueError(f"CPF “{texto}” não tem 11 dígitos")
        d = digitos.zfill(11)  # o Excel costuma comer os zeros da frente
        return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}"
    if tipo == "cep":
        digitos = _so_digitos(valor)
        if len(digitos) > 8 or not digitos:
            raise ValueError(f"CEP “{texto}” não tem 8 dígitos")
        d = digitos.zfill(8)
        return f"{d[:5]}-{d[5:]}"
    if tipo == "email":
        from django.core.exceptions import ValidationError
        from django.core.validators import validate_email
        try:
            validate_email(texto.lower())
        except ValidationError:
            raise ValueError(f"e-mail “{texto}” inválido")
        return texto.lower()
    if tipo == "uf":
        sigla = texto.upper() if texto.upper() in SIGLAS_UF else NOMES_UF.get(chave(texto))
        if not sigla:
            raise ValueError(f"UF “{texto}” não reconhecida")
        return sigla
    if tipo == "situacao":
        if chave(texto) not in SITUACOES:
            raise ValueError(f"situação “{texto}” não reconhecida (use: " + ", ".join(r for _, r in Irmao.SITUACOES) + ")")
        return SITUACOES[chave(texto)]
    if tipo == "estado_civil":
        if chave(texto) not in ESTADOS_CIVIS:
            raise ValueError(f"estado civil “{texto}” não reconhecido")
        return ESTADOS_CIVIS[chave(texto)]
    if tipo == "grau":
        if chave(texto) not in graus:
            raise ValueError(f"grau “{texto}” não existe nesta Loja (cadastre em Administração › Graus)")
        return graus[chave(texto)]
    limite = TAMANHOS.get(campo, 150)
    if len(texto) > limite:
        raise ValueError(f"{TITULOS.get(campo, campo)} com mais de {limite} letras")
    return texto


def _graus_da_loja(loja):
    """Nomes aceitos na coluna Grau -> id do grau. Aceita também 1, 2, 3 e abreviações comuns."""
    aceitos = {}
    for g in Grau.objects.filter(loja=loja):
        aceitos[chave(g.nome)] = g.pk
        aceitos.setdefault(str(g.ordem), g.pk)
        aceitos.setdefault(chave(g.nome) + "macom", g.pk)
    apelidos = {"aprendiz": ["aprm", "ap", "aprendizmacom"], "companheiro": ["compm", "comp", "companheiromacom"],
                "mestre": ["mm", "mestremacom", "m"], "mestreinstalado": ["mi", "mestreinst"]}
    for nome, lista in apelidos.items():
        if nome in aceitos:
            for apelido in lista:
                aceitos.setdefault(apelido, aceitos[nome])
    return aceitos


def ler_planilha(arquivo):
    """Devolve (títulos, linhas) da primeira aba. Cada linha é (número no Excel, [valores])."""
    try:
        from openpyxl import load_workbook
        pasta = load_workbook(io.BytesIO(arquivo.read()), read_only=True, data_only=True)
    except Exception:
        raise ErroDePlanilha("Não consegui abrir o arquivo. Salve a planilha como “Pasta de Trabalho do Excel (.xlsx)” e envie de novo.")
    aba = pasta.worksheets[0]
    cabecalho, linhas = None, []
    for numero, valores in enumerate(aba.iter_rows(values_only=True), 1):
        valores = list(valores)
        if all(v is None or _texto(v) == "" for v in valores):
            continue
        if cabecalho is None:
            # a linha de títulos é a primeira (entre as 15 iniciais) que tem uma coluna de nome
            if any(POR_CHAVE.get(chave(v)) == "nome" for v in valores):
                cabecalho = [_texto(v) for v in valores]
            elif numero > 15:
                break
            continue
        linhas.append((numero, valores))
        if len(linhas) > MAX_LINHAS:
            raise ErroDePlanilha(f"A planilha tem mais de {MAX_LINHAS} linhas. Divida em arquivos menores.")
    pasta.close()
    if cabecalho is None:
        raise ErroDePlanilha("Não encontrei a linha de títulos. A planilha precisa ter uma coluna chamada “Nome” "
                             "(ou “Nome completo”) na primeira aba. Baixe a planilha modelo para ver o formato.")
    if not linhas:
        raise ErroDePlanilha("A planilha tem os títulos, mas nenhuma linha de dados.")
    return cabecalho, linhas


def analisar(loja, arquivo):
    """Lê e confere a planilha SEM gravar nada. Devolve um resumo que pode ser guardado e depois gravado."""
    cabecalho, linhas = ler_planilha(arquivo)
    extras = {chave(e.rotulo): e for e in CampoExtra.objects.filter(loja=loja)}
    mapa, ignoradas, usados = {}, [], set()
    for posicao, titulo in enumerate(cabecalho):
        if not titulo:
            continue
        campo = POR_CHAVE.get(chave(titulo))
        if campo is None and chave(titulo) in extras:
            campo = f"extra:{extras[chave(titulo)].pk}"
        if campo is None or campo in usados:
            ignoradas.append(titulo)
            continue
        usados.add(campo)
        mapa[posicao] = campo
    graus = _graus_da_loja(loja)
    existentes = list(Irmao.objects.filter(loja=loja).values("pk", "nome", "cim", "cpf"))
    por_cim = {chave(i["cim"]): i["pk"] for i in existentes if i["cim"]}
    por_cpf = {re.sub(r"\D", "", i["cpf"]): i["pk"] for i in existentes if i["cpf"]}
    por_nome = {chave(i["nome"]): i["pk"] for i in existentes}

    registros, erros, vistos = [], [], {}
    for numero, valores in linhas:
        if any("apague esta linha" in _texto(v).lower() for v in valores):
            continue  # linha de exemplo da planilha modelo
        dados, problemas = {}, []
        for posicao, campo in mapa.items():
            valor = valores[posicao] if posicao < len(valores) else None
            if campo.startswith("extra:"):
                if _texto(valor):
                    dados[campo] = _texto(valor)[:200]
                continue
            try:
                convertido = converter(campo, TIPOS[campo], valor, graus)
            except ValueError as erro:
                problemas.append(str(erro))
                continue
            if convertido not in ("", None):
                dados[campo] = convertido.isoformat() if isinstance(convertido, date) else convertido
        if not dados.get("nome"):
            problemas.insert(0, "sem nome")
        identidades = [f"{tipo}:{valor}" for tipo, valor in (("cim", chave(dados.get("cim"))),
                                                             ("cpf", re.sub(r"\D", "", dados.get("cpf", ""))),
                                                             ("nome", chave(dados.get("nome")))) if valor]
        repetida = next((vistos[i] for i in identidades if i in vistos), None)
        if repetida:
            problemas.append(f"repetido na planilha (mesmo Irmão da linha {repetida})")
        if problemas:
            erros.append({"linha": numero, "nome": dados.get("nome", ""), "problemas": problemas})
            continue
        for identidade in identidades:
            vistos[identidade] = numero
        pk = (por_cim.get(chave(dados.get("cim"))) or por_cpf.get(re.sub(r"\D", "", dados.get("cpf", "")))
              or por_nome.get(chave(dados["nome"])))
        registros.append({"linha": numero, "pk": pk, "dados": dados})
    return {"colunas": [(cabecalho[p], TITULOS.get(c) or next(e.rotulo for e in extras.values() if f"extra:{e.pk}" == c))
                        for p, c in mapa.items()],
            "ignoradas": ignoradas, "registros": registros, "erros": erros,
            "novos": sum(1 for r in registros if not r["pk"]), "atualizados": sum(1 for r in registros if r["pk"])}


@transaction.atomic
def gravar(loja, analise):
    """Grava o que a análise aprovou. Células vazias nunca apagam o que já está cadastrado."""
    graus = {g.pk: g for g in Grau.objects.filter(loja=loja)}
    por_ordem = {}
    for g in sorted(graus.values(), key=lambda g: (g.ordem, g.pk)):
        por_ordem.setdefault(g.ordem, g)
    criados = atualizados = 0
    for registro in analise["registros"]:
        dados = dict(registro["dados"])
        irmao = Irmao.objects.filter(loja=loja, pk=registro["pk"]).first() if registro["pk"] else None
        novo = irmao is None
        if novo:
            irmao = Irmao(loja=loja)
        grau_id = dados.pop("grau", None)
        conjuge, conjuge_nasc = dados.pop("conjuge", None), dados.pop("conjuge_nascimento", None)
        datas_grau = {c: date.fromisoformat(dados[c]) for c in GRAU_POR_DATA if dados.get(c)}
        extras = dict(irmao.extras or {})
        for campo, valor in dados.items():
            if campo.startswith("extra:"):
                extras[campo.split(":")[1]] = valor
            elif campo in ("data_elevacao", "data_exaltacao", "data_instalacao"):
                continue
            elif TIPOS.get(campo) == "data":
                setattr(irmao, campo, date.fromisoformat(valor))
            else:
                setattr(irmao, campo, valor)
        irmao.extras = extras
        irmao.save()
        # Controle de graus: cada data informada vira um registro (sem duplicar)
        for campo, quando in datas_grau.items():
            ordem, descricao = GRAU_POR_DATA[campo]
            grau = por_ordem.get(ordem)
            if grau and not HistoricoGrau.objects.filter(irmao=irmao, grau=grau).exists():
                HistoricoGrau.objects.create(loja=loja, irmao=irmao, grau=grau, data=quando, observacao=descricao)
        if irmao.graus.exists():
            servicos.atualizar_grau(irmao)
            irmao.refresh_from_db(fields=["grau"])
        # A coluna Grau vale quando a planilha não trouxe datas que levem a um grau mais alto
        if grau_id in graus and (irmao.grau is None or graus[grau_id].ordem > irmao.grau.ordem):
            irmao.grau = graus[grau_id]
            irmao.save(update_fields=["grau"])
        if conjuge:
            dependente = irmao.dependentes.filter(parentesco="conjuge").first()
            if dependente is None:
                dependente = Dependente(loja=loja, irmao=irmao, parentesco="conjuge")
            dependente.nome = conjuge
            if conjuge_nasc:
                dependente.data_nascimento = date.fromisoformat(conjuge_nasc)
            dependente.save()
        criados += novo
        atualizados += not novo
    return criados, atualizados


def planilha_modelo(loja):
    """Gera a planilha modelo (.xlsx) com os títulos certos, uma linha de exemplo e as instruções."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    pasta = Workbook()
    aba = pasta.active
    aba.title = "Irmãos"
    titulos = [c[1] for c in COLUNAS] + [e.rotulo for e in CampoExtra.objects.filter(loja=loja)]
    exemplo = {"Nome completo": "João da Silva Exemplo (apague esta linha)", "CIM": "123456", "Situação": "Ativo",
               "Grau": "Mestre", "Cargo": "Orador", "Data de nascimento": date(1975, 3, 25), "CPF": "123.456.789-09",
               "E-mail": "joao@exemplo.com.br", "Celular": "(11) 98765-4321", "Endereço": "Rua das Acácias, 100",
               "Bairro": "Centro", "Cidade": "São Paulo", "UF": "SP", "CEP": "01000-000", "Estado civil": "Casado",
               "Data de iniciação": date(2010, 5, 14), "Data de elevação": date(2011, 6, 10),
               "Data de exaltação": date(2012, 8, 3), "Cônjuge": "Maria Exemplo"}
    aba.append(titulos)
    aba.append([exemplo.get(t, "") for t in titulos])
    for coluna, titulo in enumerate(titulos, 1):
        celula = aba.cell(row=1, column=coluna)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F3A5F")
        celula.alignment = Alignment(vertical="center", wrap_text=True)
        aba.column_dimensions[get_column_letter(coluna)].width = max(16, min(len(titulo) + 6, 42))
        if "Data" in titulo or "Nascimento" in titulo:
            for linha in range(2, 502):
                aba.cell(row=linha, column=coluna).number_format = "DD/MM/YYYY"
        if titulo in ("CPF", "CEP", "CIM", "RG", "Celular", "Telefone"):
            for linha in range(2, 502):
                aba.cell(row=linha, column=coluna).number_format = "@"  # texto: não perde zeros à esquerda
    aba.row_dimensions[1].height = 32
    aba.freeze_panes = "B2"
    ajuda = pasta.create_sheet("Como preencher")
    for linha in [
        ["Como preencher a planilha"], [""],
        ["1.", "Uma linha por Irmão, na aba “Irmãos”. Só a coluna “Nome completo” é obrigatória."],
        ["2.", "Não mude os títulos da primeira linha. Colunas que você não usa podem ficar vazias ou ser apagadas."],
        ["3.", "Datas no formato dia/mês/ano, por exemplo 25/03/1975."],
        ["4.", "Situação: " + ", ".join(r for _, r in Irmao.SITUACOES) + ". Em branco = Ativo."],
        ["5.", "Grau: " + ", ".join(g.nome for g in Grau.objects.filter(loja=loja)) + "."],
        ["6.", "As datas de iniciação, elevação, exaltação e instalação alimentam o Controle de Graus."],
        ["7.", "Se o Irmão já existir (mesmo CIM, CPF ou nome), o cadastro é atualizado em vez de duplicado."],
        ["8.", "Célula vazia nunca apaga um dado que já está no sistema."],
        ["9.", "Antes de gravar, o sistema mostra uma conferência com o que será incluído, atualizado ou recusado."],
    ]:
        ajuda.append(linha)
    ajuda["A1"].font = Font(bold=True, size=14)
    ajuda.column_dimensions["B"].width = 110
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()

"""Testes automáticos. Rode com:  python manage.py test"""
from io import StringIO
import json
import tempfile
from datetime import date, timedelta
from decimal import Decimal
from unittest import mock

from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from . import boleto, push, servicos
from .crud import RECURSOS
from .models import (Ata, Bem, Boleto, CampoExtra, CertificadoVisita, ContaBancaria, Emprestimo, Evento, Grau,
                     HistoricoGrau, InscricaoPush, Irmao, Lancamento, Licenca, Loja, Notificacao, PlanoConta, Presenca,
                     Sessao, Transferencia, Usuario)
from .relatorios import RELATORIOS

SENHA = "senha-de-teste-123"
D = Decimal


class Base(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.hoje = timezone.localdate()
        cls.loja = Loja.objects.create(nome="Loja A", numero="1", cnpj="00.000.000/0001-91", valor_mensalidade=D("100"))
        cls.outra = Loja.objects.create(nome="Loja B", numero="2")
        for l in (cls.loja, cls.outra):
            servicos.preparar_loja(l)
        cls.graus = {g.nome: g for g in Grau.objects.filter(loja=cls.loja)}
        cls.irmaos = [Irmao.objects.create(loja=cls.loja, nome=f"Irmão {n}", email=f"i{n}@exemplo.com.br",
                                           grau=cls.graus["Mestre"], data_nascimento=date(1970, 5, 20))
                      for n in range(4)]
        cls.alheio = Irmao.objects.create(loja=cls.outra, nome="Irmão de outra Loja", email="x@exemplo.com.br")
        novo = lambda login, **k: Usuario.objects.create_user(login, f"{login}@exemplo.com.br", SENHA, **k)  # noqa: E731
        cls.master = novo("master", loja=cls.loja, is_master=True, irmao=cls.irmaos[0])
        cls.secretario = novo("secretario", loja=cls.loja, acesso_secretaria=True)
        cls.tesoureiro = novo("tesoureiro", loja=cls.loja, acesso_tesouraria=True)
        cls.chanceler = novo("chanceler", loja=cls.loja, acesso_chancelaria=True)
        cls.hospitaleiro = novo("hospitaleiro", loja=cls.loja, acesso_hospitalaria=True)
        cls.comum = novo("comum", loja=cls.loja, irmao=cls.irmaos[1])
        cls.master_b = novo("masterb", loja=cls.outra, is_master=True)
        cls.plataforma = Usuario.objects.create_superuser("plataforma", "p@exemplo.com.br", SENHA)
        cls.caixa = ContaBancaria.objects.get(loja=cls.loja)
        cls.banco = ContaBancaria.objects.create(loja=cls.loja, nome="Banco", banco_codigo="001", banco_nome="Banco Exemplo",
                                                 agencia="1234", conta="56789", carteira="17", convenio="1234567")
        cls.plano = {p.codigo: p for p in PlanoConta.objects.filter(loja=cls.loja)}

    def entrar(self, usuario):
        self.client.force_login(usuario)

    def receber(self, irmao=None, valor="100", vencimento=None, **k):
        return Lancamento.objects.create(loja=self.loja, tipo="R", descricao="Mensalidade", irmao=irmao or self.irmaos[1],
                                         plano=self.plano["1.1"], valor=D(valor), vencimento=vencimento or self.hoje, **k)


class SeparacaoEntreLojas(Base):
    """Uma Loja nunca pode ver nem alterar dados de outra."""

    def test_listas_so_mostram_a_propria_loja(self):
        self.entrar(self.master)
        resposta = self.client.get(reverse("irmaos_lista"))
        self.assertContains(resposta, "Irmão 0")
        self.assertNotContains(resposta, "Irmão de outra Loja")

    def test_nao_abre_registro_de_outra_loja(self):
        self.entrar(self.master)
        for nome in ("irmaos_editar", "irmaos_excluir", "irmao_ficha"):
            self.assertEqual(self.client.get(reverse(nome, args=[self.alheio.pk])).status_code, 404, nome)
        self.client.post(reverse("irmaos_excluir", args=[self.alheio.pk]))
        self.assertTrue(Irmao.objects.filter(pk=self.alheio.pk).exists())

    def test_nao_aceita_ligar_registro_a_irmao_de_outra_loja(self):
        self.entrar(self.master)
        resposta = self.client.post(reverse("dependentes_novo"), {"irmao": self.alheio.pk, "nome": "Fulana", "parentesco": "conjuge"})
        self.assertEqual(resposta.status_code, 200)  # voltou para o formulário com erro
        self.assertFalse(self.alheio.dependentes.exists())
        self.assertNotContains(self.client.get(reverse("dependentes_novo")), "Irmão de outra Loja")

    def test_registro_novo_cai_na_loja_de_quem_cadastra(self):
        self.entrar(self.master_b)
        self.client.post(reverse("bens_novo"), {"descricao": "Andador", "estado": "bom"})
        self.assertEqual(Bem.objects.get(descricao="Andador").loja, self.outra)

    def test_master_nao_alcanca_usuario_de_outra_loja(self):
        self.entrar(self.master_b)
        self.assertEqual(self.client.get(reverse("usuarios_editar", args=[self.master.pk])).status_code, 404)
        self.assertNotContains(self.client.get(reverse("usuarios_lista")), "secretario")

    def test_relatorios_e_frequencia_nao_misturam_lojas(self):
        self.entrar(self.master_b)
        resposta = self.client.get(reverse("relatorio", args=["quadro"]))
        self.assertContains(resposta, "Irmão de outra Loja")
        self.assertNotContains(resposta, "Irmão 0")

    def test_boleto_e_arquivo_de_outra_loja(self):
        l = self.receber()
        b = boleto.emitir(l, self.banco)
        self.entrar(self.master_b)
        self.assertEqual(self.client.get(reverse("boleto_ver", args=[b.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("arquivo", args=[f"loja_{self.loja.pk}/atas/x.pdf"])).status_code, 404)

    def test_loja_bloqueada_derruba_o_acesso(self):
        self.entrar(self.master_b)
        Loja.objects.filter(pk=self.outra.pk).update(ativa=False)
        self.assertRedirects(self.client.get(reverse("inicio")), reverse("entrar"), fetch_redirect_response=False)

    def test_mesmo_nome_de_grau_em_lojas_diferentes(self):
        self.entrar(self.master)
        resposta = self.client.post(reverse("graus_novo"), {"nome": "Mestre", "ordem": 3})
        self.assertContains(resposta, "já existe")  # repetido dentro da mesma Loja: recusa com aviso, sem quebrar
        self.assertEqual(Grau.objects.filter(nome="Mestre").count(), 2)


class Acessos(Base):
    def test_cada_cargo_so_entra_no_seu_modulo(self):
        portas = {"secretaria": "irmaos_lista", "tesouraria": "receber_lista", "chancelaria": "sessoes_lista",
                  "hospitalaria": "bens_lista", "admin": "usuarios_lista"}
        casos = {self.secretario: {"secretaria"}, self.tesoureiro: {"tesouraria"}, self.chanceler: {"chancelaria"},
                 self.hospitaleiro: {"hospitalaria"}, self.comum: set(), self.master: set(portas)}
        for usuario, liberados in casos.items():
            self.entrar(usuario)
            for modulo, tela in portas.items():
                esperado = 200 if modulo in liberados else 403
                self.assertEqual(self.client.get(reverse(tela)).status_code, esperado, f"{usuario} em {modulo}")

    def test_sem_login_vai_para_a_entrada(self):
        resposta = self.client.get(reverse("irmaos_lista"))
        self.assertEqual(resposta.status_code, 302)
        self.assertIn(reverse("entrar"), resposta.url)

    def test_irmao_comum_ve_os_proprios_dados_e_so_eles(self):
        meu, de_outro = self.receber(self.irmaos[1]), self.receber(self.irmaos[2], valor="77.77")
        meu_boleto, boleto_de_outro = boleto.emitir(meu, self.banco), boleto.emitir(de_outro, self.banco)
        self.entrar(self.comum)
        pagina = self.client.get(reverse("minha_area"))
        self.assertContains(pagina, "Irmão 1")
        self.assertNotContains(pagina, "77,77")
        self.assertEqual(self.client.get(reverse("boleto_ver", args=[meu_boleto.pk])).status_code, 200)
        self.assertEqual(self.client.get(reverse("boleto_ver", args=[boleto_de_outro.pk])).status_code, 403)
        self.assertEqual(self.client.get(reverse("relatorio", args=["inadimplencia"])).status_code, 403)

    def test_master_nao_consegue_virar_administrador_da_plataforma(self):
        self.entrar(self.master)
        self.client.post(reverse("usuarios_editar", args=[self.comum.pk]), {
            "username": "comum", "is_active": "on", "is_superuser": "on", "is_staff": "on", "loja": self.outra.pk})
        self.comum.refresh_from_db()
        self.assertFalse(self.comum.is_superuser)
        self.assertEqual(self.comum.loja, self.loja)
        self.assertEqual(self.client.get(reverse("plataforma_lojas")).status_code, 403)

    def test_master_cria_usuario_com_acesso_de_cargo(self):
        self.entrar(self.master)
        self.client.post(reverse("usuarios_novo"), {"username": "novo.tes", "senha": "uma-senha-boa-99", "is_active": "on",
                                                    "acesso_tesouraria": "on", "irmao": self.irmaos[3].pk})
        novo = Usuario.objects.get(username="novo.tes")
        self.assertEqual((novo.loja, novo.irmao, novo.pode("tesouraria"), novo.pode("secretaria")),
                         (self.loja, self.irmaos[3], True, False))
        self.assertTrue(novo.check_password("uma-senha-boa-99"))

    def test_criar_acesso_para_todos(self):
        self.entrar(self.master)
        self.client.post(reverse("usuarios_criar_todos"))
        self.assertEqual(Usuario.objects.filter(loja=self.loja, irmao__isnull=False).count(), 4)
        self.assertFalse(Usuario.objects.filter(irmao=self.alheio).exists())

    def test_plataforma_cria_loja_e_entra_nela(self):
        self.entrar(self.plataforma)
        self.assertRedirects(self.client.get(reverse("inicio")), reverse("plataforma_lojas"))
        self.client.post(reverse("plataforma_nova_loja"), {"nome": "Loja C", "master_login": "masterc", "master_nome": "João Silva",
                                                           "master_email": "c@exemplo.com.br", "master_senha": "senha-forte-321"})
        nova = Loja.objects.get(nome="Loja C")
        self.assertEqual(Grau.objects.filter(loja=nova).count(), 4)
        self.assertTrue(Usuario.objects.get(username="masterc").is_master)
        self.client.post(reverse("plataforma_lojas"), {"loja": nova.pk, "acao": "entrar"})
        self.assertContains(self.client.get(reverse("inicio")), "Loja C")

    def test_personalizacao_muda_cores_e_campo_extra_aparece_no_cadastro(self):
        self.entrar(self.master)
        dados = {"nome": "Loja A Renomeada", "cor_primaria": "#7A1F2B", "cor_destaque": "#C9A227", "valor_mensalidade": "120",
                 "dia_vencimento": 5, "frequencia_minima": 60}
        self.assertRedirects(self.client.post(reverse("personalizar"), dados), reverse("personalizar"))
        self.assertContains(self.client.get(reverse("inicio")), "--primaria:#7A1F2B")
        self.client.post(reverse("campos-extras_novo"), {"rotulo": "Nome simbólico", "ordem": 1})
        extra = CampoExtra.objects.get(loja=self.loja)
        self.assertContains(self.client.get(reverse("irmaos_novo")), "Nome simbólico")
        self.client.post(reverse("irmaos_novo"), {"nome": "Novo Irmão", "situacao": "ativo", "receber_emails": "on",
                                                  f"extra_{extra.pk}": "Hiram"})
        self.assertEqual(Irmao.objects.get(nome="Novo Irmão").extras, {str(extra.pk): "Hiram"})
        self.entrar(self.master_b)
        self.assertNotContains(self.client.get(reverse("irmaos_novo")), "Nome simbólico")


class TodasAsTelasAbrem(Base):
    def test_cadastros(self):
        self.entrar(self.master)
        for slug in RECURSOS:
            for sufixo in ("_lista", "_novo"):
                self.assertEqual(self.client.get(reverse(slug + sufixo)).status_code, 200, slug + sufixo)

    def test_demais_telas(self):
        sessao = Sessao.objects.create(loja=self.loja, data=self.hoje)
        servicos.sincronizar_presencas(sessao)
        ata = Ata.objects.create(loja=self.loja, data=self.hoje, titulo="Ata 1", texto="Texto")
        l = self.receber()
        self.entrar(self.master)
        telas = ["inicio", "minha_area", "emails", "mensalidades_serie", "boletos", "relatorios", "notificacoes", "personalizar",
                 "usuarios_criar_todos", "senha_trocar", "sw", "manifesto",
                 ("irmao_ficha", self.irmaos[0].pk), ("ata_ver", ata.pk), ("sessao_presenca", sessao.pk),
                 ("lancamento_quitar", l.pk), ("lancamento_cancelar", l.pk), ("boleto_gerar", l.pk)]
        for tela in telas:
            url = reverse(tela) if isinstance(tela, str) else reverse(tela[0], args=[tela[1]])
            self.assertEqual(self.client.get(url).status_code, 200, url)
        self.client.logout()
        for tela in ("entrar", "senha_esqueci"):
            self.assertEqual(self.client.get(reverse(tela)).status_code, 200)

    def test_todos_os_relatorios_na_tela_e_em_planilha(self):
        sessao = Sessao.objects.create(loja=self.loja, data=self.hoje)
        servicos.sincronizar_presencas(sessao)
        servicos.quitar(self.receber(), self.hoje, D("100"), self.caixa)
        self.receber(vencimento=self.hoje - timedelta(days=40))
        boleto.emitir(self.receber(), self.banco)
        Emprestimo.objects.create(loja=self.loja, bem=Bem.objects.create(loja=self.loja, descricao="Cadeira"), beneficiario="Fulano")
        self.entrar(self.master)
        self.assertGreaterEqual(len(RELATORIOS), 25)
        for slug in RELATORIOS:
            url = reverse("relatorio", args=[slug])
            self.assertEqual(self.client.get(url).status_code, 200, slug)
            planilha = self.client.get(url, {"formato": "csv"})
            self.assertEqual(planilha.status_code, 200, slug)
            self.assertIn("text/csv", planilha["Content-Type"])


class Tesouraria(Base):
    def test_saldo_quitacao_estorno_e_transferencia(self):
        l = self.receber(valor="150")
        despesa = Lancamento.objects.create(loja=self.loja, tipo="P", descricao="Luz", favorecido="Cia", plano=self.plano["2.3"],
                                            valor=D("40"), vencimento=self.hoje)
        self.assertEqual(servicos.saldo_conta(self.caixa), D("0"))
        servicos.quitar(l, self.hoje, D("150"), self.caixa)
        servicos.quitar(despesa, self.hoje, D("40"), self.caixa)
        Transferencia.objects.create(loja=self.loja, origem=self.caixa, destino=self.banco, valor=D("60"), data=self.hoje)
        self.assertEqual(servicos.saldo_conta(self.caixa), D("50"))
        self.assertEqual(servicos.saldo_conta(self.banco), D("60"))
        servicos.estornar(l)
        self.assertEqual(servicos.saldo_conta(self.caixa), D("-100"))
        l.refresh_from_db()
        self.assertEqual((l.status, l.valor_pago), ("aberto", None))

    def test_telas_de_quitar_estornar_e_cancelar(self):
        l = self.receber(valor="100")
        self.entrar(self.tesoureiro)
        self.client.post(reverse("lancamento_quitar", args=[l.pk]), {"data": self.hoje.isoformat(), "valor": "95.50", "conta": self.caixa.pk})
        l.refresh_from_db()
        self.assertEqual((l.status, l.valor_pago, l.conta), ("pago", D("95.50"), self.caixa))
        self.client.post(reverse("lancamento_estornar", args=[l.pk]))
        self.client.post(reverse("lancamento_cancelar", args=[l.pk]))
        l.refresh_from_db()
        self.assertEqual(l.status, "cancelado")

    def test_lancamento_exige_plano_do_tipo_certo(self):
        self.entrar(self.tesoureiro)
        base = {"descricao": "Teste", "favorecido": "Alguém", "valor": "10", "vencimento": self.hoje.isoformat()}
        self.client.post(reverse("receber_novo"), {**base, "plano": self.plano["2.3"].pk})  # despesa em conta a receber
        self.client.post(reverse("receber_novo"), {**base, "plano": self.plano["1"].pk})  # conta que só agrupa
        self.assertFalse(Lancamento.objects.exists())
        self.client.post(reverse("pagar_novo"), {**base, "plano": self.plano["2.3"].pk})
        self.assertEqual(Lancamento.objects.get().tipo, "P")

    def test_mensalidades_em_serie_nao_duplicam(self):
        self.entrar(self.tesoureiro)
        dados = {"plano": self.plano["1.1"].pk, "descricao": "Mensalidade", "valor": "100", "meses": 12,
                 "primeiro_vencimento": "2026-01-31", "irmaos": [i.pk for i in self.irmaos[:3]]}
        self.client.post(reverse("mensalidades_serie"), dados)
        self.assertEqual(Lancamento.objects.count(), 36)
        vencimentos = sorted(set(Lancamento.objects.values_list("vencimento", flat=True)))
        self.assertEqual(vencimentos[1], date(2026, 2, 28))  # dia 31 vira o último dia de fevereiro
        self.assertEqual(vencimentos[2], date(2026, 3, 31))
        self.client.post(reverse("mensalidades_serie"), {**dados, "irmaos": [i.pk for i in self.irmaos]})
        self.assertEqual(Lancamento.objects.count(), 48)  # só o 4º Irmão ganhou as 12 dele

    def test_balancete_fecha_com_os_saldos(self):
        self.caixa.saldo_inicial = D("1000")
        self.caixa.save()
        mes_passado = self.hoje.replace(day=1) - timedelta(days=1)
        servicos.quitar(self.receber(valor="300"), mes_passado, D("300"), self.caixa)
        servicos.quitar(self.receber(valor="200"), self.hoje, D("200"), self.caixa)
        despesa = Lancamento.objects.create(loja=self.loja, tipo="P", descricao="Luz", favorecido="Cia", plano=self.plano["2.3"],
                                            valor=D("50"), vencimento=self.hoje)
        servicos.quitar(despesa, self.hoje, D("50"), self.caixa)
        Transferencia.objects.create(loja=self.loja, origem=self.caixa, destino=self.banco, valor=D("100"), data=self.hoje)
        b = servicos.balancete(self.loja, self.hoje.year, self.hoje.month)
        self.assertEqual((b["total_receitas"], b["total_despesas"], b["resultado"]), (D("200"), D("50"), D("150")))
        self.assertEqual((b["saldo_anterior"], b["saldo_final"]), (D("1300"), D("1450")))
        self.assertEqual(b["saldo_final"], b["saldo_anterior"] + b["resultado"])  # transferência não muda o total
        self.assertEqual(b["saldo_final"], servicos.saldo_conta(self.caixa) + servicos.saldo_conta(self.banco))


class Boletos(Base):
    def test_digitos_verificadores_com_exemplo_publicado(self):
        # Exemplo clássico da documentação de cobrança do Banco do Brasil.
        barras = "00193373700000001000500940144816060680935031"
        self.assertEqual(boleto.modulo11(barras[:4] + barras[5:]), barras[4])
        self.assertEqual(boleto.linha_digitavel(barras), "00190.50095 40144.816069 06809.350314 3 37370000000100")

    def test_fator_de_vencimento_inclusive_a_virada_de_2025(self):
        self.assertEqual(boleto.fator_vencimento(date(2000, 7, 3)), "1000")
        self.assertEqual(boleto.fator_vencimento(date(2025, 2, 21)), "9999")
        self.assertEqual(boleto.fator_vencimento(date(2025, 2, 22)), "1000")
        self.assertEqual(boleto.fator_vencimento(date(2026, 10, 11)), "1596")

    def test_boleto_tem_44_digitos_coerentes_e_nosso_numero_sequencial(self):
        b1, b2 = boleto.emitir(self.receber(valor="180.50"), self.banco), boleto.emitir(self.receber(), self.banco)
        self.assertEqual(len(b1.codigo_barras), 44)
        self.assertEqual(b1.codigo_barras[:4], "0019")
        self.assertEqual(b1.codigo_barras[9:19], "0000018050")
        self.assertEqual(boleto.modulo11(b1.codigo_barras[:4] + b1.codigo_barras[5:]), b1.codigo_barras[4])
        self.assertEqual(int(b2.nosso_numero), int(b1.nosso_numero) + 1)
        self.assertEqual(boleto.emitir(b1.lancamento, self.banco), b1)  # não gera dois boletos para a mesma conta
        self.assertEqual(sum(w for _, w in boleto.barras_para_desenho(b1.codigo_barras)), 4 + 22 * 18 + 5)

    def _retorno_do_banco(self, remessa, movimento="06", pago=None):
        """Simula o arquivo que o banco devolveria para os boletos de uma remessa."""
        linhas = []
        for linha in remessa.conteudo.splitlines():
            if linha[7] == "3" and linha[13] == "P":
                t = list(" " * 240)
                t[0:8], t[13], t[15:17] = linha[0:7] + "3", "T", movimento
                t[37:57], t[81:96], t[105:130] = linha[37:57], linha[85:100], linha[195:220]
                u = list(" " * 240)
                u[0:8], u[13], u[15:17] = linha[0:7] + "3", "U", movimento
                u[77:92] = pago or linha[85:100]
                u[137:145] = self.hoje.strftime("%d%m%Y")
                linhas += ["".join(t), "".join(u)]
        return "\r\n".join(linhas)

    def test_remessa_cnab_e_baixa_automatica_pelo_retorno(self):
        lancamentos = [self.receber(self.irmaos[n], valor="180") for n in range(3)]
        self.entrar(self.tesoureiro)
        self.client.post(reverse("boletos"), {"acao": "lote", "conta": self.banco.pk})
        self.assertEqual(Boleto.objects.count(), 3)
        self.client.post(reverse("boletos"), {"acao": "remessa", "conta": self.banco.pk})
        remessa = self.loja and Boleto.objects.first().remessa
        linhas = remessa.conteudo.splitlines()
        self.assertEqual(len(linhas), 2 + 3 * 2 + 2)
        self.assertTrue(all(len(l) == 240 for l in linhas))
        self.assertEqual([l[7] for l in linhas], ["0", "1", "3", "3", "3", "3", "3", "3", "5", "9"])
        self.assertEqual(linhas[2][85:100], "000000000018000")
        self.assertEqual(linhas[-1][23:29], "000010")  # total de registros do arquivo
        self.assertEqual(set(Boleto.objects.values_list("status", flat=True)), {"remetido"})
        baixado = self.client.get(reverse("remessa_baixar", args=[remessa.pk]))
        self.assertIn("attachment", baixado["Content-Disposition"])
        # segunda remessa não reenvia os mesmos boletos
        self.assertIsNone(boleto.gerar_remessa(self.loja, self.banco))

        arquivo = SimpleUploadedFile("retorno.ret", self._retorno_do_banco(remessa).encode("latin-1"))
        self.client.post(reverse("boletos"), {"acao": "retorno", "arquivo": arquivo})
        for l in lancamentos:
            l.refresh_from_db()
            self.assertEqual((l.status, l.valor_pago, l.data_pagamento, l.conta), ("pago", D("180"), self.hoje, self.banco))
        self.assertEqual(set(Boleto.objects.values_list("status", flat=True)), {"pago"})
        self.assertEqual(servicos.saldo_conta(self.banco), D("540"))
        # o mesmo retorno enviado duas vezes não quita em dobro
        boleto.processar_retorno(self.loja, "de novo.ret", self._retorno_do_banco(remessa))
        self.assertEqual(servicos.saldo_conta(self.banco), D("540"))

    def test_retorno_de_registro_e_de_outra_loja(self):
        b = boleto.emitir(self.receber(), self.banco)
        remessa = boleto.gerar_remessa(self.loja, self.banco)
        boleto.processar_retorno(self.loja, "r.ret", self._retorno_do_banco(remessa, movimento="02"))
        b.refresh_from_db()
        self.assertEqual((b.status, b.lancamento.status), ("registrado", "aberto"))
        r = boleto.processar_retorno(self.outra, "r.ret", self._retorno_do_banco(remessa))  # outra Loja sobe o arquivo
        b.lancamento.refresh_from_db()
        self.assertEqual((r.quitados, r.nao_encontrados, b.lancamento.status), (0, 1, "aberto"))


class Chancelaria(Base):
    def _sessoes(self, n, **k):
        sessoes = [Sessao.objects.create(loja=self.loja, data=self.hoje - timedelta(days=7 * i), **k) for i in range(n)]
        for s in sessoes:
            servicos.sincronizar_presencas(s)
        return sessoes

    def _marcar(self, irmao, situacoes):
        for p, status in zip(Presenca.objects.filter(irmao=irmao).order_by("sessao__data"), situacoes):
            p.status = status
            p.save()

    def test_presenca_com_um_clique(self):
        sessao = self._sessoes(1)[0]
        p = sessao.presencas.get(irmao=self.irmaos[0])
        self.entrar(self.chanceler)
        r = self.client.post(reverse("presenca_marcar", args=[p.pk]), {"status": "P"})
        self.assertEqual(r.json()["resumo"], {"P": 1, "F": 3, "A": 0, "L": 0})
        r = self.client.post(reverse("presenca_marcar", args=[p.pk]), {"status": "A", "justificativa": "Viagem"})
        p.refresh_from_db()
        self.assertEqual((p.status, p.justificativa), ("A", "Viagem"))
        self.assertEqual(self.client.post(reverse("presenca_marcar", args=[p.pk]), {"status": "X"}).status_code, 400)
        self.client.post(reverse("sessao_presenca", args=[sessao.pk]), {"acao": "todos"})
        self.assertEqual(sessao.presencas.filter(status="P").count(), 3)  # o abonado continua abonado
        self.entrar(self.secretario)
        self.assertEqual(self.client.post(reverse("presenca_marcar", args=[p.pk]), {"status": "P"}).status_code, 403)

    def test_criar_sessao_pela_tela_monta_a_chamada(self):
        Licenca.objects.create(loja=self.loja, irmao=self.irmaos[3], inicio=self.hoje - timedelta(days=1), fim=self.hoje + timedelta(days=1))
        self.irmaos[2].situacao = "desligado"
        self.irmaos[2].save()
        self.entrar(self.chanceler)
        self.client.post(reverse("sessoes_novo"), {"data": self.hoje.isoformat(), "tipo": "ordinaria"})
        situacoes = dict(Presenca.objects.values_list("irmao__nome", "status"))
        self.assertEqual(situacoes, {"Irmão 0": "F", "Irmão 1": "F", "Irmão 3": "L"})

    def test_sessao_de_grau_so_chama_quem_tem_o_grau(self):
        self.irmaos[0].grau = self.graus["Aprendiz"]
        self.irmaos[0].save()
        sessao = self._sessoes(1, grau=self.graus["Mestre"])[0]
        self.assertNotIn(self.irmaos[0].pk, sessao.presencas.values_list("irmao_id", flat=True))
        self.assertEqual(sessao.presencas.count(), 3)

    def test_calculo_de_frequencia(self):
        self._sessoes(10)
        a, b, c, d = self.irmaos
        self._marcar(a, "PPPPPPFFFF")  # 60%
        self._marcar(b, "PPPPAAFFFF")  # 4 presenças + 2 abonos = 60%
        self._marcar(c, "PPPPLLLLLF")  # 5 sessões fora por licença: 4 de 5 = 80%
        self._marcar(d, "PPPPFFFFFF")  # 40% + 2 visitas = 60%
        for n in (5, 20):
            CertificadoVisita.objects.create(loja=self.loja, irmao=d, data=self.hoje - timedelta(days=n), loja_visitada="Loja X")
        inicio = self.hoje - timedelta(days=100)

        def pct():
            return {f["irmao"].nome: (f["percentual"], f["apto"]) for f in servicos.frequencia(self.loja, inicio, self.hoje)}
        self.assertEqual(pct(), {"Irmão 0": (60.0, True), "Irmão 1": (60.0, True), "Irmão 2": (80.0, True), "Irmão 3": (60.0, True)})
        Loja.objects.filter(pk=self.loja.pk).update(abono_conta_presenca=False, visita_conta_presenca=False, frequencia_minima=60)
        self.loja.refresh_from_db()
        self.assertEqual(pct(), {"Irmão 0": (60.0, True), "Irmão 1": (40.0, False), "Irmão 2": (80.0, True), "Irmão 3": (40.0, False)})
        # fora do período não conta
        ultima_semana = servicos.frequencia(self.loja, self.hoje - timedelta(days=6), self.hoje, a)[0]
        self.assertEqual(ultima_semana["sessoes"], 1)

    def test_visitas_nao_passam_de_cem_por_cento(self):
        self._sessoes(2)
        self._marcar(self.irmaos[0], "PF")
        for n in range(5):
            CertificadoVisita.objects.create(loja=self.loja, irmao=self.irmaos[0], data=self.hoje - timedelta(days=n), loja_visitada="X")
        self.assertEqual(servicos.frequencia(self.loja, self.hoje - timedelta(days=30), self.hoje, self.irmaos[0])[0]["percentual"], 100.0)

    def test_relatorio_de_frequencia_mostra_aptos(self):
        self._sessoes(4)
        self._marcar(self.irmaos[0], "PPPP")
        self.entrar(self.chanceler)
        pagina = self.client.get(reverse("relatorio", args=["frequencia"])).content.decode()
        self.assertIn("1 apto(s) de 4", pagina)
        self.assertIn("100,0%", pagina)


class SecretariaEHospitalaria(Base):
    def test_controle_de_graus_atualiza_o_grau_atual(self):
        irmao = Irmao.objects.create(loja=self.loja, nome="Recém iniciado")
        self.entrar(self.secretario)
        for grau, quando in (("Aprendiz", "2024-03-01"), ("Companheiro", "2025-03-01")):
            self.client.post(reverse("graus-historico_novo"), {"irmao": irmao.pk, "grau": self.graus[grau].pk, "data": quando})
        irmao.refresh_from_db()
        self.assertEqual(irmao.grau, self.graus["Companheiro"])
        self.assertEqual(HistoricoGrau.objects.filter(irmao=irmao).count(), 2)

    @override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
    def test_disparo_de_email_individual_e_respeita_quem_nao_quer_receber(self):
        Irmao.objects.filter(pk=self.irmaos[3].pk).update(receber_emails=False)
        self.entrar(self.secretario)
        self.client.post(reverse("emails"), {"publico": "ativos", "assunto": "Convocação", "corpo": "Prezado Ir∴ {nome}, compareça."})
        self.assertEqual(len(mail.outbox), 3)
        self.assertTrue(all(len(m.to) == 1 for m in mail.outbox))
        self.assertEqual({m.to[0] for m in mail.outbox}, {"i0@exemplo.com.br", "i1@exemplo.com.br", "i2@exemplo.com.br"})
        self.assertIn("Prezado Ir∴ Irmão 0", next(m.body for m in mail.outbox if m.to == ["i0@exemplo.com.br"]))
        self.assertContains(self.client.get(reverse("emails")), "Convocação")

    def test_ata_com_arquivo_fica_protegida(self):
        with tempfile.TemporaryDirectory() as pasta, self.settings(MEDIA_ROOT=__import__("pathlib").Path(pasta)):
            self.entrar(self.secretario)
            pdf = SimpleUploadedFile("ata.pdf", b"%PDF-1.4 teste", content_type="application/pdf")
            self.client.post(reverse("atas_novo"), {"data": self.hoje.isoformat(), "titulo": "Ata 10", "arquivo": pdf})
            ata = Ata.objects.get()
            self.assertTrue(ata.arquivo.name.startswith(f"loja_{self.loja.pk}/atas/"))
            url = reverse("arquivo", args=[ata.arquivo.name])
            self.assertEqual(b"".join(self.client.get(url).streaming_content), b"%PDF-1.4 teste")
            self.entrar(self.comum)  # Irmão sem cargo na Secretaria não baixa ata
            self.assertEqual(self.client.get(url).status_code, 403)
            self.entrar(self.master_b)  # outra Loja nem fica sabendo que existe
            self.assertEqual(self.client.get(url).status_code, 404)
            self.entrar(self.secretario)  # arquivo perigoso é recusado
            self.client.post(reverse("atas_novo"), {"data": self.hoje.isoformat(), "titulo": "Ata ruim",
                                                    "arquivo": SimpleUploadedFile("x.html", b"<script>")})
            self.assertFalse(Ata.objects.filter(titulo="Ata ruim").exists())

    def test_emprestimo_e_devolucao(self):
        self.entrar(self.hospitaleiro)
        self.client.post(reverse("bens_novo"), {"descricao": "Cadeira de rodas", "estado": "bom", "codigo": "H-1"})
        bem = Bem.objects.get()
        dados = {"bem": bem.pk, "beneficiario": "Dona Maria", "data_emprestimo": self.hoje.isoformat()}
        self.client.post(reverse("emprestimos_novo"), dados)
        self.assertEqual(Emprestimo.objects.count(), 1)
        self.client.post(reverse("emprestimos_novo"), {**dados, "beneficiario": "Seu José"})  # já está emprestada
        self.assertEqual(Emprestimo.objects.count(), 1)
        self.assertContains(self.client.get(reverse("bens_lista")), "Emprestado a Dona Maria")
        self.client.post(reverse("emprestimo_devolver", args=[Emprestimo.objects.get().pk]))
        self.assertEqual(Emprestimo.objects.get().data_devolucao, self.hoje)
        self.client.post(reverse("emprestimos_novo"), {**dados, "beneficiario": "Seu José"})
        self.assertEqual(Emprestimo.objects.count(), 2)
        self.client.post(reverse("bens_excluir", args=[bem.pk]))  # bem com histórico não pode sumir
        self.assertTrue(Bem.objects.filter(pk=bem.pk).exists())


CHAVES = push.gerar_chaves()


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend", PUSH_CHAVE_PUBLICA=CHAVES[0],
                   PUSH_CHAVE_PRIVADA=CHAVES[1])
class AvisosAutomaticos(Base):
    def test_rotina_diaria_aniversario_e_lembrete_sem_repetir(self):
        Irmao.objects.filter(pk=self.irmaos[1].pk).update(data_nascimento=date(1980, self.hoje.month, self.hoje.day))
        Sessao.objects.create(loja=self.loja, data=self.hoje, tipo="magna")
        from datetime import datetime, time
        agora = timezone.make_aware(datetime.combine(self.hoje, time(12, 0)))  # meio-dia de hoje, no horário local
        Evento.objects.create(loja=self.loja, titulo="Jantar", inicio=agora, local="Salão")
        Evento.objects.create(loja=self.loja, titulo="Semana que vem", inicio=agora + timedelta(days=7))
        resumo = servicos.rotina_diaria(self.loja)
        self.assertEqual(resumo, "1 aniversariante(s), 1 e-mail(s), 2 lembrete(s)")
        self.assertEqual(mail.outbox[0].to, ["i1@exemplo.com.br"])
        self.assertIn("Feliz aniversário, Ir∴ Irmão 1", mail.outbox[0].subject)
        do_master = list(self.master.notificacoes.values_list("texto", flat=True))
        self.assertEqual(len(do_master), 3)
        self.assertTrue(any("aniversário do Ir∴ Irmão 1" in t for t in do_master))
        self.assertEqual(self.comum.notificacoes.count(), 2)  # o aniversariante não é avisado do próprio aniversário
        self.assertFalse(Notificacao.objects.filter(loja=self.outra).exists())
        self.assertIsNone(servicos.rotina_diaria(self.loja))  # rodar de novo no mesmo dia não repete nada
        self.assertEqual(len(mail.outbox), 1)

    def test_aniversario_de_29_de_fevereiro(self):
        self.assertTrue(servicos.faz_aniversario(date(1980, 2, 29), date(2027, 2, 28)))
        self.assertFalse(servicos.faz_aniversario(date(1980, 2, 29), date(2028, 2, 28)))
        self.assertTrue(servicos.faz_aniversario(date(1980, 2, 29), date(2028, 2, 29)))

    def test_assinatura_do_toque_e_valida(self):
        from cryptography.hazmat.primitives import hashes
        from cryptography.hazmat.primitives.asymmetric import ec
        from cryptography.hazmat.primitives.asymmetric.utils import encode_dss_signature
        cabecalho = push.cabecalho_vapid("https://fcm.googleapis.com/fcm/send/abc")
        token = cabecalho.split("t=")[1].split(",")[0]
        partes = token.split(".")
        corpo = json.loads(push._b64_ler(partes[1]))
        self.assertEqual(corpo["aud"], "https://fcm.googleapis.com")
        assinatura = push._b64_ler(partes[2])
        publica = ec.EllipticCurvePublicKey.from_encoded_point(ec.SECP256R1(), push._b64_ler(CHAVES[0]))
        publica.verify(encode_dss_signature(int.from_bytes(assinatura[:32], "big"), int.from_bytes(assinatura[32:], "big")),
                       f"{partes[0]}.{partes[1]}".encode(), ec.ECDSA(hashes.SHA256()))  # levanta erro se inválida

    def test_inscricao_toque_e_busca_dos_avisos(self):
        self.entrar(self.comum)
        enviar = lambda endpoint: self.client.post(reverse("push_inscrever"), json.dumps({"endpoint": endpoint}),  # noqa: E731
                                                   content_type="application/json").status_code
        self.assertEqual(enviar("http://169.254.169.254/latest/meta-data"), 400)
        self.assertEqual(enviar("https://servidor-interno.local/x"), 400)
        self.assertEqual(enviar("https://fcm.googleapis.com.atacante.com/x"), 400)
        self.assertEqual(enviar("https://fcm.googleapis.com/fcm/send/abc"), 200)
        self.assertEqual(enviar("https://updates.push.services.mozilla.com/wpush/v2/xyz"), 200)
        with mock.patch("gestao.push._enviar", side_effect=[201, 410]) as toque:
            aceitos = servicos.notificar(self.loja, [self.comum, self.master], "Lembrete: é hoje!", "Sessão Ordinária")
        self.assertEqual((aceitos, toque.call_count), (1, 2))
        self.assertEqual(InscricaoPush.objects.count(), 1)  # a inscrição cancelada (410) foi removida
        avisos = self.client.get(reverse("push_pendentes")).json()["avisos"]
        self.assertEqual(avisos, [{"titulo": "Lembrete: é hoje!", "texto": "Sessão Ordinária"}])
        self.assertEqual(self.client.get(reverse("push_pendentes")).json()["avisos"], [])  # não mostra duas vezes
        self.assertContains(self.client.get(reverse("notificacoes")), "Sessão Ordinária")


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
class TabelaDeAcessos(TestCase):
    def _rodar(self, linhas, *extras):
        from io import StringIO
        from django.core.management import call_command
        with tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False, encoding="utf-8") as f:
            f.write("perfil;login;nome;email;loja\n" + "\n".join(linhas))
        saida = StringIO()
        call_command("criar_acessos", f.name, *extras, stdout=saida)
        return saida.getvalue()

    def test_cria_administrador_com_senha_e_os_demais_pelo_email(self):
        from django.core.management.base import CommandError
        tabela = ["administrador;administrador;Administrador;;", "master;veneravel;João Silva;v@exemplo.com.br;Loja Nova",
                  "tesouraria+secretaria;tes;Pedro Lima;t@exemplo.com.br;Loja Nova", "irmao;ir;Zé;z@exemplo.com.br;Loja Nova"]
        saida = self._rodar(tabela, "--senha-admin", "Senha.Forte.9431", "--site", "sistema.exemplo.com.br", "--enviar-emails")
        self.assertNotIn("Senha.Forte.9431", saida)  # a senha nunca aparece na tela
        admin = Usuario.objects.get(username="administrador")
        self.assertTrue(admin.is_superuser and admin.check_password("Senha.Forte.9431") and admin.loja is None)
        loja = Loja.objects.get(nome="Loja Nova")
        self.assertEqual(Grau.objects.filter(loja=loja).count(), 4)
        v, t, i = (Usuario.objects.get(username=u) for u in ("veneravel", "tes", "ir"))
        self.assertEqual((v.loja, v.is_master, v.is_superuser), (loja, True, False))
        self.assertEqual((t.pode("tesouraria"), t.pode("secretaria"), t.pode("chancelaria")), (True, True, False))
        self.assertEqual(i.modulos, [])
        self.assertEqual(sorted(m.to[0] for m in mail.outbox), ["t@exemplo.com.br", "v@exemplo.com.br", "z@exemplo.com.br"])
        link = next(p for p in mail.outbox[0].body.split() if p.startswith("https://sistema.exemplo.com.br/senha/nova/"))
        caminho = link.replace("https://sistema.exemplo.com.br", "")
        pagina = self.client.get(caminho, follow=True)
        self.client.post(pagina.redirect_chain[-1][0], {"new_password1": "minha-senha-nova-77", "new_password2": "minha-senha-nova-77"})
        self.assertTrue(self.client.login(username=mail.outbox[0].to[0][0] == "v" and "veneravel" or
                                          Usuario.objects.get(email=mail.outbox[0].to[0]).username, password="minha-senha-nova-77"))
        # rodar de novo não troca a senha de quem já criou a sua nem duplica ninguém
        self._rodar(tabela, "--senha-admin", "Senha.Forte.9431")
        self.assertEqual(Usuario.objects.count(), 4)
        self.assertTrue(Usuario.objects.get(email=mail.outbox[0].to[0]).check_password("minha-senha-nova-77"))
        with self.assertRaises(CommandError):
            self._rodar(["secretaria;sec;Sem Email;;Loja Nova"])
        with self.assertRaises(CommandError):
            self._rodar(["administrador;adm2;Outro;;"], "--senha-admin", "123")


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
class SenhaPorEmail(Base):
    DADOS = {"username": "novo.irmao", "first_name": "Carlos", "email": "carlos@exemplo.com.br", "is_active": "on",
             "senha": "Inicial-2026-xy", "enviar_senha": "on"}

    def test_senha_cadastrada_vai_para_o_email_e_precisa_ser_trocada(self):
        self.entrar(self.master)
        self.client.post(reverse("usuarios_novo"), self.DADOS)
        novo = Usuario.objects.get(username="novo.irmao")
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ["carlos@exemplo.com.br"])
        self.assertIn("Nome de acesso: novo.irmao", mail.outbox[0].body)
        self.assertIn("Senha inicial: Inicial-2026-xy", mail.outbox[0].body)
        self.assertNotIn("Inicial-2026-xy", novo.password)  # no banco só fica o embaralhado, nunca a senha
        self.client.logout()
        self.assertTrue(self.client.login(username="novo.irmao", password="Inicial-2026-xy"))
        self.assertRedirects(self.client.get(reverse("minha_area")), reverse("senha_trocar"))
        self.client.post(reverse("senha_trocar"), {"old_password": "Inicial-2026-xy", "new_password1": "so-minha-senha-88",
                                                   "new_password2": "so-minha-senha-88"})
        self.assertEqual(self.client.get(reverse("minha_area")).status_code, 200)
        self.assertFalse(Usuario.objects.get(username="novo.irmao").deve_trocar_senha)

    def test_sem_marcar_o_envio_ou_sem_email(self):
        self.entrar(self.master)
        self.client.post(reverse("usuarios_novo"), {**self.DADOS, "enviar_senha": ""})
        self.assertEqual(len(mail.outbox), 0)
        self.assertFalse(Usuario.objects.get(username="novo.irmao").deve_trocar_senha)
        resposta = self.client.post(reverse("usuarios_novo"), {**self.DADOS, "username": "outro", "email": ""})
        self.assertContains(resposta, "Preencha o e-mail para enviar a senha")
        self.assertFalse(Usuario.objects.filter(username="outro").exists())
        # editar sem digitar senha não envia nada nem mexe na senha
        u = Usuario.objects.get(username="novo.irmao")
        self.client.post(reverse("usuarios_editar", args=[u.pk]), {**self.DADOS, "senha": "", "first_name": "Carlos Alberto"})
        u.refresh_from_db()
        self.assertEqual((u.first_name, u.check_password("Inicial-2026-xy"), len(mail.outbox)), ("Carlos Alberto", True, 0))


class AcessoDoAdministrador(TestCase):
    def test_demo_cria_o_administrador_com_a_senha_propria(self):
        from io import StringIO
        from django.core.management import call_command
        from .acesso_admin import ADMIN_LOGIN, ADMIN_SENHA_HASH
        call_command("demo", stdout=StringIO())
        admin = Usuario.objects.get(is_superuser=True)
        self.assertEqual((admin.username, admin.password), (ADMIN_LOGIN, ADMIN_SENHA_HASH))
        self.assertFalse(admin.check_password("loja-demo-2026"))  # não usa a senha de exemplo
        self.assertFalse(Usuario.objects.filter(username__in=["plataforma", "cristi286913ano"]).exists())
        self.client.force_login(admin)
        self.assertEqual(self.client.get(reverse("plataforma_lojas")).status_code, 200)


class GarantirAdministrador(TestCase):
    def test_cria_quando_falta_e_nao_mexe_na_senha_de_quem_ja_trocou(self):
        from io import StringIO
        from django.core.management import call_command
        from .acesso_admin import ADMIN_LOGIN, ADMIN_SENHA_HASH
        call_command("garantir_admin", stdout=StringIO())
        admin = Usuario.objects.get()
        self.assertEqual((admin.username, admin.password, admin.is_superuser), (ADMIN_LOGIN, ADMIN_SENHA_HASH, True))
        admin.set_password("outra-senha-minha-55")
        admin.is_superuser = False
        admin.last_login = timezone.now()
        admin.save()
        call_command("garantir_admin", stdout=StringIO())
        admin.refresh_from_db()
        self.assertTrue(admin.is_superuser and admin.check_password("outra-senha-minha-55"))
        self.assertEqual(Usuario.objects.count(), 1)
        call_command("garantir_admin", "--redefinir", stdout=StringIO())
        self.assertEqual(Usuario.objects.get().password, ADMIN_SENHA_HASH)


class AdministradorQueNuncaEntrou(TestCase):
    def test_senha_errada_e_corrigida_ate_o_primeiro_acesso(self):
        from io import StringIO
        from django.core.management import call_command
        from django.utils import timezone as tz
        from .acesso_admin import ADMIN_LOGIN, ADMIN_SENHA_HASH
        admin = Usuario.objects.create_superuser(ADMIN_LOGIN, "", "senha-antiga-errada-1")
        call_command("garantir_admin", stdout=StringIO())
        admin.refresh_from_db()
        self.assertEqual(admin.password, ADMIN_SENHA_HASH)  # nunca entrou: volta para a senha original
        admin.set_password("troquei-a-minha-senha-9")
        admin.last_login = tz.now()
        admin.save()
        call_command("garantir_admin", stdout=StringIO())
        admin.refresh_from_db()
        self.assertTrue(admin.check_password("troquei-a-minha-senha-9"))  # já entrou: a instalação não mexe mais
        self.assertContains(self.client.get(reverse("entrar")), "Versão 19/10/2026-C")


def importacao_colunas():
    from .importacao import COLUNAS
    return COLUNAS


class ImportarPlanilha(Base):
    def _xlsx(self, linhas, antes=()):
        import io
        from openpyxl import Workbook
        pasta = Workbook()
        for linha in [*antes, *linhas]:
            pasta.active.append(linha)
        saida = io.BytesIO()
        pasta.save(saida)
        return SimpleUploadedFile("cadastro.xlsx", saida.getvalue())

    def _enviar(self, linhas, **k):
        return self.client.post(reverse("importar_irmaos"), {"arquivo": self._xlsx(linhas, **k)}, follow=True)

    def test_importa_confere_antes_e_grava_so_na_confirmacao(self):
        from datetime import datetime
        self.entrar(self.secretario)
        extra = CampoExtra.objects.create(loja=self.loja, rotulo="Padrinho")
        pagina = self._enviar([
            ["NOME", "Nº CIM", "Dt. Nascimento", "cpf", "E-mail", "Estado", "Grau", "Iniciação", "Exaltação", "Esposa",
             "Padrinho", "Time de futebol", "Situação", "CEP"],
            ["José Novo da Silva", 4455, datetime(1975, 3, 25), 1234567809, "JOSE@Exemplo.com.br", "São Paulo", "Mestre",
             "14/05/2010", "03/08/2012", "Maria Novo", "Irmão 0", "Santos", "regular", 1000000],
            ["Irmão 1", None, "20/05/1971", None, None, "RJ", None, None, None, None, None, None, None, None],
            ["Linha Ruim", None, "31/02/1980", None, "sem-arroba", None, "Grau 33", None, None, None, None, None, None, None],
            [None, 99, None, None, None, None, None, None, None, None, None, None, None, None],
            ["José Novo da Silva", 4455, None, None, None, None, None, None, None, None, None, None, None, None],
        ], antes=[["Cadastro geral da Loja"], [None]])
        self.assertContains(pagina, "Time de futebol")  # coluna ignorada é avisada
        self.assertContains(pagina, "não reconhecida")
        self.assertContains(pagina, "repetido na planilha")
        analise = self.client.session["importacao"]["analise"]
        self.assertEqual((analise["novos"], analise["atualizados"], len(analise["erros"])), (1, 1, 3))
        self.assertFalse(Irmao.objects.filter(nome="José Novo da Silva").exists())  # ainda não gravou nada
        self.client.post(reverse("importar_irmaos"), {"acao": "gravar"})
        novo = Irmao.objects.get(nome="José Novo da Silva")
        self.assertEqual((novo.loja, novo.cim, novo.cpf, novo.email, novo.uf, novo.cep, novo.situacao),
                         (self.loja, "4455", "012.345.678-09", "jose@exemplo.com.br", "SP", "01000-000", "ativo"))
        self.assertEqual((novo.data_nascimento, novo.data_iniciacao), (date(1975, 3, 25), date(2010, 5, 14)))
        self.assertEqual(novo.grau, self.graus["Mestre"])
        self.assertEqual(sorted(novo.graus.values_list("grau__nome", "data")),
                         [("Aprendiz", date(2010, 5, 14)), ("Mestre", date(2012, 8, 3))])
        self.assertEqual(novo.dependentes.get().nome, "Maria Novo")
        self.assertEqual(novo.extras, {str(extra.pk): "Irmão 0"})
        existente = Irmao.objects.get(pk=self.irmaos[1].pk)
        self.assertEqual((existente.data_nascimento, existente.uf, existente.email),
                         (date(1971, 5, 20), "RJ", "i1@exemplo.com.br"))  # atualizou; célula vazia não apagou o e-mail
        self.assertEqual(Irmao.objects.filter(loja=self.loja).count(), 5)
        self.assertFalse(Irmao.objects.filter(nome="Linha Ruim").exists())
        self.assertNotIn("importacao", self.client.session)
        # enviar a mesma planilha de novo não duplica ninguém
        self._enviar([["Nome", "CIM"], ["José Novo da Silva", 4455]])
        self.client.post(reverse("importar_irmaos"), {"acao": "gravar"})
        self.assertEqual(Irmao.objects.filter(loja=self.loja).count(), 5)

    def test_arquivos_invalidos_modelo_e_acessos(self):
        self.entrar(self.secretario)
        r = self.client.post(reverse("importar_irmaos"), {"arquivo": SimpleUploadedFile("x.xlsx", b"nao e excel")}, follow=True)
        self.assertContains(r, "Não consegui abrir o arquivo")
        r = self.client.post(reverse("importar_irmaos"), {"arquivo": SimpleUploadedFile("x.csv", b"a;b")}, follow=True)
        self.assertContains(r, ".xlsx")
        self.assertContains(self._enviar([["Coluna A", "Coluna B"], ["x", "y"]]), "Não encontrei a linha de títulos")
        self.assertContains(self.client.post(reverse("importar_irmaos"), {"acao": "gravar"}, follow=True), "expirou")
        # a planilha modelo é reconhecida por inteiro, e a linha de exemplo dela nunca é importada
        modelo = self.client.get(reverse("importar_modelo"))
        self.assertIn("spreadsheetml", modelo["Content-Type"])
        r = self.client.post(reverse("importar_irmaos"), {"arquivo": SimpleUploadedFile("modelo.xlsx", modelo.content)}, follow=True)
        analise = self.client.session["importacao"]["analise"]
        self.assertEqual((analise["novos"], analise["erros"], analise["ignoradas"]), (0, [], []))
        self.assertEqual(len(analise["colunas"]), len(importacao_colunas()))
        self.client.post(reverse("importar_irmaos"), {"acao": "cancelar"})
        self.assertNotIn("importacao", self.client.session)
        # só a Secretaria importa, e a conferência de uma Loja não vale em outra
        self.entrar(self.tesoureiro)
        self.assertEqual(self.client.get(reverse("importar_irmaos")).status_code, 403)
        self.entrar(self.secretario)
        self._enviar([["Nome"], ["Só da Loja A"]])
        sessao = self.client.session
        sessao["importacao"] = {**sessao["importacao"], "loja": self.outra.pk}
        sessao.save()
        self.client.post(reverse("importar_irmaos"), {"acao": "gravar"})
        self.assertFalse(Irmao.objects.filter(nome="Só da Loja A").exists())


class ExcluirVarios(Base):
    def test_exclui_selecionados_so_com_a_palavra_e_respeita_tesouraria_e_loja(self):
        url = reverse("irmaos_excluir_varios")
        self.receber(self.irmaos[1])  # Irmão 1 tem lançamento na Tesouraria
        from .models import Dependente
        Dependente.objects.create(loja=self.loja, irmao=self.irmaos[0], nome="Dependente X")
        self.entrar(self.secretario)
        self.assertEqual(self.client.get(url).status_code, 405)  # só por envio da tela, nunca por link
        ids = [self.irmaos[0].pk, self.irmaos[1].pk, self.alheio.pk]
        pagina = self.client.post(url, {"ids": ids})
        self.assertContains(pagina, "2 Irmão(s)")  # o de outra Loja nem entra na conta
        self.assertNotContains(pagina, "Irmão de outra Loja")
        self.assertNotContains(pagina, 'name="financeiro"')  # Secretaria sozinha não apaga financeiro
        self.assertEqual(Irmao.objects.filter(loja=self.loja).count(), 4)
        self.assertContains(self.client.post(url, {"ids": ids, "confirmar": "1", "palavra": "sim"}), "digite a palavra EXCLUIR")
        self.assertEqual(Irmao.objects.filter(loja=self.loja).count(), 4)
        self.client.post(url, {"ids": ids, "confirmar": "1", "palavra": "excluir", "financeiro": "1"})
        self.assertEqual(sorted(Irmao.objects.filter(loja=self.loja).values_list("nome", flat=True)), ["Irmão 1", "Irmão 2", "Irmão 3"])
        self.assertFalse(Dependente.objects.exists())
        self.assertTrue(Irmao.objects.filter(pk=self.alheio.pk).exists())
        self.assertEqual(Lancamento.objects.count(), 1)

    def test_master_exclui_todos_inclusive_financeiro_e_filtro_vale_para_todas_as_paginas(self):
        url = reverse("irmaos_excluir_varios")
        b = boleto.emitir(self.receber(self.irmaos[1]), self.banco)
        Irmao.objects.filter(pk=self.irmaos[3].pk).update(situacao="desligado")
        self.entrar(self.master)
        pagina = self.client.post(url, {"todos": "1", "situacao": "desligado"})
        self.assertContains(pagina, "1 Irmão(s)")  # "todos" respeita o filtro que estava na tela
        pagina = self.client.post(url, {"todos": "1"})
        self.assertContains(pagina, "4 Irmão(s)")
        self.assertContains(pagina, 'name="financeiro"')
        ids = [i.pk for i in self.irmaos]
        self.client.post(url, {"ids": ids, "confirmar": "1", "palavra": "EXCLUIR"})  # sem marcar o financeiro
        self.assertEqual(list(Irmao.objects.filter(loja=self.loja).values_list("nome", flat=True)), ["Irmão 1"])
        self.master.refresh_from_db()
        self.assertTrue(self.master.is_active and self.master.irmao is None)  # o usuário continua, só perde a ligação
        self.client.post(url, {"ids": ids, "confirmar": "1", "palavra": "EXCLUIR", "financeiro": "1"})
        self.assertFalse(Irmao.objects.filter(loja=self.loja).exists())
        self.assertFalse(Lancamento.objects.exists() or Boleto.objects.filter(pk=b.pk).exists())
        self.assertContains(self.client.get(reverse("irmaos_lista")), "Ainda não há nada aqui")
        self.entrar(self.tesoureiro)
        self.assertEqual(self.client.post(url, {"todos": "1"}).status_code, 403)


class ExcluirVariosNosModulos(Base):
    def test_toda_lista_de_modulo_tem_selecao_menos_a_de_usuarios(self):
        self.entrar(self.master)
        for slug, recurso in RECURSOS.items():
            tem = recurso.selecao is not None
            self.assertEqual(tem, recurso.modulo != "admin", slug)
        self.receber()
        self.assertContains(self.client.get(reverse("receber_lista"), {"ver": "lista"}), 'id="excluir-selecionados"')
        self.assertNotContains(self.client.get(reverse("usuarios_lista")), 'id="excluir-selecionados"')
        with self.assertRaises(Exception):
            reverse("usuarios_excluir_varios")

    def test_tesouraria_exclui_lancamentos_com_aviso_palavra_e_sem_tocar_outra_loja(self):
        url = reverse("receber_excluir_varios")
        aberto, pago, a_pagar = self.receber(), self.receber(valor="50"), None
        servicos.quitar(pago, self.hoje, D("50"), self.caixa)
        a_pagar = Lancamento.objects.create(loja=self.loja, tipo="P", descricao="Luz", favorecido="Cia", plano=self.plano["2.3"],
                                            valor=D("40"), vencimento=self.hoje)
        plano_b = PlanoConta.objects.filter(loja=self.outra, codigo="1.1").get()
        alheio = Lancamento.objects.create(loja=self.outra, tipo="R", descricao="De outra Loja", favorecido="X", plano=plano_b,
                                           valor=D("10"), vencimento=self.hoje)
        ids = [aberto.pk, pago.pk, a_pagar.pk, alheio.pk]
        self.entrar(self.secretario)
        self.assertEqual(self.client.post(url, {"ids": ids}).status_code, 403)
        self.entrar(self.tesoureiro)
        pagina = self.client.post(url, {"ids": ids})
        self.assertContains(pagina, "Excluir 2 registro(s)")  # só os "a receber" desta Loja
        self.assertContains(pagina, "1 já estão quitados")
        self.client.post(url, {"ids": ids, "confirmar": "1", "palavra": "nao"})
        self.assertEqual(Lancamento.objects.count(), 4)
        self.assertEqual(servicos.saldo_conta(self.caixa), D("50"))
        self.client.post(url, {"ids": ids, "confirmar": "1", "palavra": "EXCLUIR"})
        self.assertEqual(sorted(Lancamento.objects.values_list("descricao", flat=True)), ["De outra Loja", "Luz"])
        self.assertEqual(servicos.saldo_conta(self.caixa), D("0"))
        self.assertRedirects(self.client.get(url), reverse("receber_lista"))

    def test_plano_de_contas_apaga_o_que_pode_e_mantem_o_que_esta_em_uso(self):
        self.receber()  # usa a conta 1.1
        self.entrar(self.tesoureiro)
        antes = PlanoConta.objects.filter(loja=self.loja).count()
        self.client.post(reverse("plano-contas_excluir_varios"), {"todos": "1", "confirmar": "0"})
        ids = list(PlanoConta.objects.filter(loja=self.loja).values_list("pk", flat=True))
        resposta = self.client.post(reverse("plano-contas_excluir_varios"), {"ids": ids, "confirmar": "1", "palavra": "EXCLUIR"}, follow=True)
        self.assertContains(resposta, "foram mantidos")
        # sobram só a conta em uso (1.1) e a conta que a agrupa (1); filhas são apagadas antes das mães
        self.assertEqual(sorted(PlanoConta.objects.filter(loja=self.loja).values_list("codigo", flat=True)), ["1", "1.1"])
        self.assertEqual(PlanoConta.objects.filter(loja=self.outra).count(), antes)

    def test_chancelaria_exclui_sessoes_com_as_presencas(self):
        sessoes = [Sessao.objects.create(loja=self.loja, data=self.hoje - timedelta(days=7 * n)) for n in range(3)]
        for s in sessoes:
            servicos.sincronizar_presencas(s)
        Licenca.objects.create(loja=self.loja, irmao=self.irmaos[0], inicio=self.hoje, fim=self.hoje)
        self.entrar(self.chanceler)
        url = reverse("sessoes_excluir_varios")
        pagina = self.client.post(url, {"todos": "1"})
        self.assertContains(pagina, "Excluir 3 registro(s)")
        self.assertContains(pagina, "12 marcações de presença")
        self.client.post(url, {"ids": [sessoes[0].pk, sessoes[1].pk], "confirmar": "1", "palavra": "excluir"})
        self.assertEqual((Sessao.objects.count(), Presenca.objects.count()), (1, 4))
        self.client.post(reverse("licencas_excluir_varios"), {"todos": "1", "confirmar": "0"})
        self.client.post(reverse("licencas_excluir_varios"), {"ids": list(Licenca.objects.values_list("pk", flat=True)),
                                                              "confirmar": "1", "palavra": "EXCLUIR"})
        self.assertFalse(Licenca.objects.exists())
        self.entrar(self.tesoureiro)
        self.assertEqual(self.client.post(url, {"todos": "1"}).status_code, 403)


class ComandosDeImportacao(Base):
    def _planilha(self, pasta, linhas):
        from openpyxl import Workbook
        pasta_xlsx = Workbook()
        for linha in linhas:
            pasta_xlsx.active.append(linha)
        caminho = f"{pasta}/cadastro.xlsx"
        pasta_xlsx.save(caminho)
        return caminho

    def test_importar_irmaos_confere_e_so_grava_com_a_opcao(self):
        from io import StringIO
        from django.core.management import call_command
        from django.core.management.base import CommandError
        with tempfile.TemporaryDirectory() as pasta:
            caminho = self._planilha(pasta, [["Nome completo", "CIM", "Grau"], ["Irmão Planilha", 777, "Mestre"]])
            saida = StringIO()
            call_command("importar_irmaos", caminho, "--loja", "1", stdout=saida)
            self.assertIn("1 novo(s)", saida.getvalue())
            self.assertFalse(Irmao.objects.filter(nome="Irmão Planilha").exists())
            call_command("importar_irmaos", caminho, "--loja", "Loja A", "--gravar", stdout=StringIO())
            self.assertEqual(Irmao.objects.get(nome="Irmão Planilha").loja, self.loja)
            with self.assertRaises(CommandError):
                call_command("importar_irmaos", caminho, "--loja", "Loja Inexistente", stdout=StringIO())

    def test_carga_inicial_importa_uma_unica_vez(self):
        from io import StringIO
        from pathlib import Path
        from django.core.management import call_command
        with tempfile.TemporaryDirectory() as pasta:
            self._planilha(pasta, [["Nome completo", "CIM", "Celular"], ["Irmão Carga", 888, "(11) 90000-0000"]])
            saida = StringIO()
            call_command("carga_inicial", "--pasta", pasta, stdout=saida, stderr=saida)
            self.assertIn("falta o arquivo carga/loja.txt", saida.getvalue())
            Path(pasta, "loja.txt").write_text("2\n", encoding="utf-8")  # número da Loja B
            call_command("carga_inicial", "--pasta", pasta, stdout=StringIO())
            irmao = Irmao.objects.get(nome="Irmão Carga")
            self.assertEqual(irmao.loja, self.outra)
            irmao.celular = "(11) 95555-5555"  # alteração feita depois, pelo site
            irmao.save()
            saida = StringIO()
            call_command("carga_inicial", "--pasta", pasta, stdout=saida)  # próxima instalação
            self.assertIn("já tinha sido importada", saida.getvalue())
            irmao.refresh_from_db()
            self.assertEqual(irmao.celular, "(11) 95555-5555")
            self.assertEqual(Irmao.objects.filter(nome="Irmão Carga").count(), 1)


class RenomearLoja(TestCase):
    def test_migracao_renomeia_a_loja_123(self):
        from importlib import import_module
        from django.apps import apps as registro
        migracao = import_module("gestao.migrations.0006_loja_laquila_romana")
        Loja.objects.create(nome="A∴R∴L∴S∴ Fraternidade e Luz", numero="123", potencia="Grande Oriente Exemplo", oriente="X")
        Loja.objects.create(nome="Cavaleiros do Sul", numero="456")
        migracao.renomear(registro, None)
        migracao.renomear(registro, None)  # rodar de novo não muda nada
        loja = Loja.objects.get(numero="3365")
        self.assertEqual((loja.nome, loja.potencia, loja.oriente, loja.uf, str(loja)),
                         ("L'AQUILA ROMANA", "Grande Oriente do Brasil", "São Paulo", "SP", "L'AQUILA ROMANA nº 3365"))
        self.assertTrue(Loja.objects.filter(nome="Cavaleiros do Sul", numero="456").exists())


class ImportarFinanceiro(Base):
    def setUp(self):
        for irmao, nome in zip(self.irmaos, ["Alfredo Bastos Lima", "Bruno Costa", "Irmão Dois", "Daniel Reis"]):
            Irmao.objects.filter(pk=irmao.pk).update(nome=nome)
        Irmao.objects.filter(pk=self.irmaos[3].pk).update(cim="5555")

    def _xlsx(self, abas):
        import io
        from openpyxl import Workbook
        pasta = Workbook()
        pasta.remove(pasta.active)
        for titulo, linhas in abas.items():
            aba = pasta.create_sheet(titulo)
            for linha in linhas:
                aba.append(linha)
        saida = io.BytesIO()
        pasta.save(saida)
        return SimpleUploadedFile("controle.xlsx", saida.getvalue())

    def _planilha(self):
        meses = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]
        return self._xlsx({
            "Mensalidade": [
                ["Centavos ", "Nome do Obreiro", "Cotizante", "R$", *meses],
                [0.02, "  ALFREDO BASTOS LIMA   3078 ", "SIM", 180, "PAGO", "PAGO", "Acordo", None, None, None, None, None, None, None, None, None],
                [0.03, "  BRUNO COSTA", "NÂO", 60, "XXX-XXX", "PAGO", None, None, None, None, None, None, None, None, None, None],
                [0.04, "  IRMAO DOIS", "SIM", "REMIDO", "REMIDO", "REMIDO", None, None, None, None, None, None, None, None, None, None],
                [0.05, " Obreiro Novo da Silva (2309)", "SIM", 180, "PAGO", "talvez", None, None, None, None, None, None, None, None, None, None],
                [0.06, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None],
            ],
            "Tronco": [["Tronco de beneficencia"], [None], ["janeiro", None, 150.5], ["fevereiro", None, 0], ["março", None, 99],
                       ["abril", None, 0], ["maio", None, 0], ["junho", None, 0], ["julho", None, 0], ["agosot", None, 0],
                       ["setembro", None, 0], ["outubro", None, 0], ["novembro", None, 0], ["dezembro", None, 0]],
            "Livre": [["Relatório de Centavos"], [None], ["NUMERO", "CIM", "NOME", "CENTAVOS"], [1, 5555, "DANIEL REIS", 0.09], [2, 6666, "Ninguém Daqui", 0]],
            "Dados": [["texto qualquer"]],
        })

    def test_importa_verifica_e_salva_sem_duplicar(self):
        self.entrar(self.tesoureiro)
        url = reverse("importar_financeiro")
        dados = {"ano": 2026, "em_branco": "ignorar", "conta": self.caixa.pk, "incluir_novos": "on", "importar_tronco": "on"}
        pagina = self.client.post(url, {**dados, "arquivo": self._planilha()}, follow=True)
        self.assertContains(pagina, "Nada foi salvo ainda")
        self.assertContains(pagina, "“talvez” não reconhecido")
        analise = self.client.session["importacao_fin"]["analise"]
        self.assertEqual(analise["abas_ignoradas"], ["Dados"])
        self.assertEqual(analise["contagem"], {"pago": 4, "acordo": 1, "nao_se_aplica": 1, "isento": 2, "novos": 1})
        self.assertFalse(Lancamento.objects.exists())  # nada gravado na verificação
        self.client.post(url, {"acao": "gravar"})
        i0, i1, i2 = (Irmao.objects.get(pk=i.pk) for i in self.irmaos[:3])
        self.assertEqual((i0.cotizante, i0.mensalidade_individual, i0.centavos), (True, Decimal("180.00"), Decimal("0.02")))  # Loja cobra 100
        self.assertEqual((i1.cotizante, i1.mensalidade_individual), (False, Decimal("60.00")))
        self.assertEqual(i2.situacao, "remido")
        self.assertEqual(Irmao.objects.get(pk=self.irmaos[3].pk).centavos, Decimal("0.09"))  # veio da aba de centavos (CIM)
        novo = Irmao.objects.get(nome="Obreiro Novo da Silva")
        self.assertEqual((novo.loja, novo.centavos), (self.loja, Decimal("0.05")))
        mens = Lancamento.objects.filter(plano=self.plano["1.1"])
        self.assertEqual(mens.filter(status="pago").count(), 4)
        self.assertEqual(mens.filter(irmao=self.irmaos[1], status="pago").get().valor_pago, Decimal("60.00"))
        acordo = mens.get(status="aberto")
        self.assertEqual((acordo.irmao, acordo.competencia), (self.irmaos[0], date(2026, 3, 1)))
        self.assertEqual(servicos.saldo_conta(self.caixa), Decimal("180") * 3 + Decimal("60") + Decimal("150.50") + Decimal("99"))
        self.assertEqual(Lancamento.objects.filter(plano=self.plano["1.3"]).count(), 2)
        self.loja.refresh_from_db()
        self.assertEqual(self.loja.valor_mensalidade, Decimal("100"))  # a Loja já tinha valor padrão: não muda
        # mesma planilha de novo: nada duplica
        total = Lancamento.objects.count()
        self.client.post(url, {**dados, "arquivo": self._planilha()})
        self.client.post(url, {"acao": "gravar"})
        self.assertEqual(Lancamento.objects.count(), total)
        self.assertEqual(Irmao.objects.filter(nome="Obreiro Novo da Silva").count(), 1)

    def test_meses_em_branco_em_aberto_e_planilha_da_loja_volta_igual(self):
        from openpyxl import load_workbook
        import io
        self.entrar(self.tesoureiro)
        url = reverse("importar_financeiro")
        hoje = timezone.localdate()
        self.client.post(url, {"ano": hoje.year, "em_branco": "aberto", "conta": self.caixa.pk, "arquivo": self._planilha()})
        analise = self.client.session["importacao_fin"]["analise"]
        self.assertEqual(analise["ate_mes"], hoje.month)
        self.assertGreater(analise["contagem"]["aberto"], 0)
        self.client.post(url, {"acao": "gravar"})
        self.assertFalse(Irmao.objects.filter(nome="Obreiro Novo da Silva").exists())  # "incluir novos" desmarcado
        baixada = self.client.get(reverse("planilha_financeira"), {"ano": hoje.year})
        pasta = load_workbook(io.BytesIO(baixada.content))
        self.assertEqual(pasta.sheetnames, ["Mensalidade", "Centavos", "Tronco", "Como preencher"])
        linha = next(r for r in pasta["Mensalidade"].iter_rows(values_only=True) if r[1] == "Alfredo Bastos Lima")
        self.assertEqual(linha[:7], (0.02, "Alfredo Bastos Lima", None, "SIM", 180, "PAGO", "PAGO"))
        self.assertEqual(linha[7], "Acordo")
        total = Lancamento.objects.count()
        self.client.post(url, {"ano": hoje.year, "em_branco": "ignorar", "conta": self.caixa.pk,
                               "arquivo": SimpleUploadedFile("volta.xlsx", baixada.content)})
        self.client.post(url, {"acao": "gravar"})
        self.assertEqual(Lancamento.objects.count(), total)  # a planilha baixada volta sem duplicar nada
        from django.db.models import Count as _Count
        self.assertFalse(Lancamento.objects.filter(plano=self.plano["1.1"]).values("irmao", "competencia")
                         .annotate(n=_Count("id")).filter(n__gt=1).exists())
        self.assertContains(self.client.get(reverse("relatorio", args=["mensalidades-ano"]), {"ano": hoje.year}), "Acordo")
        self.assertContains(self.client.get(reverse("relatorio", args=["centavos"])), "LIVRE")

    def test_arquivo_errado_e_acesso(self):
        self.entrar(self.tesoureiro)
        url = reverse("importar_financeiro")
        base = {"ano": 2026, "em_branco": "ignorar", "conta": self.caixa.pk}
        r = self.client.post(url, {**base, "arquivo": self._xlsx({"Aba": [["Coluna A"], ["x"]]})}, follow=True)
        self.assertContains(r, "Não encontrei nenhuma aba no formato esperado")
        r = self.client.post(url, {**base, "arquivo": SimpleUploadedFile("x.csv", b"a")})
        self.assertContains(r, ".xlsx")
        self.entrar(self.secretario)
        self.assertEqual(self.client.get(url).status_code, 403)


class OrdemAlfabetica(Base):
    def test_todas_as_listas_tem_a_opcao_e_ordenam(self):
        self.entrar(self.master)
        for slug, recurso in RECURSOS.items():
            pagina = self.client.get(reverse(f"{slug}_lista"), {"ordem": "az", "ver": "lista"})
            self.assertEqual(pagina.status_code, 200, slug)
            self.assertContains(pagina, "Organizar por", msg_prefix=slug)
        Irmao.objects.filter(pk=self.irmaos[2].pk).update(nome="Álvaro Zero")
        Irmao.objects.filter(pk=self.irmaos[3].pk).update(nome="Zacarias")
        nomes = lambda ordem: [l["obj"].nome for l in self.client.get(reverse("irmaos_lista"), {"ordem": ordem}).context["linhas"]]  # noqa: E731
        az = nomes("az")
        self.assertEqual(az[-1], "Zacarias")
        self.assertEqual(nomes("za"), list(reversed(az)))

    def test_receber_alfabetica_pelo_nome_de_quem_paga(self):
        self.receber(self.irmaos[3])
        Lancamento.objects.create(loja=self.loja, tipo="R", descricao="Aluguel do salão", favorecido="Aaron Locações",
                                  plano=self.plano["1.9"], valor=D("10"), vencimento=self.hoje)
        self.entrar(self.tesoureiro)
        linhas = self.client.get(reverse("receber_lista"), {"ver": "lista", "ordem": "az"}).context["linhas"]
        self.assertEqual([l["obj"].quem for l in linhas], ["Aaron Locações", "Irmão 3"])


class ReceberPorAnoEFiltro(Base):
    def test_resumo_por_irmao_no_ano_e_filtros(self):
        hoje = self.hoje
        ano_passado = hoje.replace(year=hoje.year - 1, day=1)
        vencido = self.receber(self.irmaos[0], valor="100", vencimento=hoje.replace(day=1) if hoje.day > 1 else hoje)
        Lancamento.objects.filter(pk=vencido.pk).update(vencimento=date(hoje.year, 1, 10))
        self.receber(self.irmaos[0], valor="50", vencimento=date(hoje.year, 12, 31))
        self.receber(self.irmaos[1], valor="70", vencimento=ano_passado)
        pago = self.receber(self.irmaos[2], valor="30", vencimento=date(hoje.year, 2, 5))
        servicos.quitar(pago, date(hoje.year, 2, 5), D("30"), self.caixa)
        self.entrar(self.tesoureiro)
        url = reverse("receber_lista")
        r = self.client.get(url)  # padrão: ano atual, em aberto, resumo por Irmão
        resumo = {g["nome"]: g for g in r.context["resumo"]}
        self.assertEqual(set(resumo), {"Irmão 0"})
        self.assertEqual((resumo["Irmão 0"]["qtd"], resumo["Irmão 0"]["total"], resumo["Irmão 0"]["vencidos"]), (2, D("150"), 1))
        painel = {rotulo: valor for rotulo, valor, nota in r.context["painel"]}
        self.assertEqual((painel["Em aberto"], painel["Vencido"], painel["A vencer"], painel["Recebido"]),
                         (D("150"), D("100"), D("50"), D("30")))
        self.assertNotContains(r, 'name="ids"')  # não é mais uma lista corrida
        r = self.client.get(url, {"ano": str(hoje.year - 1), "situacao": "aberto", "ver": "irmao"})
        self.assertEqual([g["nome"] for g in r.context["resumo"]], ["Irmão 1"])
        r = self.client.get(url, {"ano": "todos", "situacao": "vencido", "ver": "lista"})
        self.assertEqual(sorted(l["obj"].valor for l in r.context["linhas"]), [D("70"), D("100")])
        r = self.client.get(url, {"ano": str(hoje.year), "situacao": "pago", "ver": "irmao"})
        self.assertEqual([g["nome"] for g in r.context["resumo"]], ["Irmão 2"])
        r = self.client.get(url, {"ano": str(hoje.year), "mes": "12", "situacao": "todos", "ver": "lista"})
        self.assertEqual([l["obj"].valor for l in r.context["linhas"]], [D("50")])
        link = self.client.get(url).context["resumo"][0]["link"]
        r = self.client.get(url + link)
        self.assertEqual({l["obj"].irmao for l in r.context["linhas"]}, {self.irmaos[0]})
        # "selecionar todos" para excluir respeita o ano e a situação da tela
        pagina = self.client.post(reverse("receber_excluir_varios"), {"todos": "1", "ano": str(hoje.year - 1), "situacao": "aberto"})
        self.assertContains(pagina, "Excluir 1 registro(s)")


class MensalidadesComRemidos(Base):
    def test_remido_fica_isento_e_situacao_e_atualizada(self):
        Irmao.objects.filter(pk=self.irmaos[3].pk).update(situacao="remido")
        self.entrar(self.tesoureiro)
        url = reverse("mensalidades_serie")
        pagina = self.client.get(url)
        self.assertContains(pagina, "Remidos (isentos de mensalidade)")
        self.assertEqual(pagina.context["form"].initial["remidos"], [self.irmaos[3].pk])
        dados = {"plano": self.plano["1.1"].pk, "descricao": "Mensalidade", "valor": "100", "meses": 3,
                 "primeiro_vencimento": "2026-11-10", "irmaos": [i.pk for i in self.irmaos],  # marcou todos, até os remidos
                 "remidos": [self.irmaos[0].pk]}  # Irmão 0 vira remido; Irmão 3 foi desmarcado
        self.client.post(url, dados)
        cobrados = set(Lancamento.objects.values_list("irmao__nome", flat=True))
        self.assertEqual(cobrados, {"Irmão 1", "Irmão 2", "Irmão 3"})
        self.assertEqual(Lancamento.objects.count(), 9)
        self.assertEqual(Irmao.objects.get(pk=self.irmaos[0].pk).situacao, "remido")
        self.assertEqual(Irmao.objects.get(pk=self.irmaos[3].pk).situacao, "ativo")
        r = self.client.post(url, {**dados, "irmaos": [self.irmaos[0].pk]})  # só um remido marcado para pagar
        self.assertContains(r, "Marque pelo menos um Irmão que paga mensalidade")

    def test_servico_nunca_cobra_remido(self):
        Irmao.objects.filter(pk=self.irmaos[0].pk).update(situacao="remido")
        criados, _ = servicos.gerar_mensalidades(self.loja, list(Irmao.objects.filter(loja=self.loja)), self.plano["1.1"],
                                                 D("100"), date(2026, 11, 10), 2)
        self.assertEqual(criados, 6)
        self.assertFalse(Lancamento.objects.filter(irmao=self.irmaos[0]).exists())


class PlanoRemidoIsento(Base):
    def test_plano_remido_lanca_isento_sem_valor_e_cancela_aberto(self):
        remido = PlanoConta.objects.get(loja=self.loja, codigo="1.6")
        self.assertTrue(remido.eh_remido)
        Irmao.objects.filter(pk=self.irmaos[0].pk).update(situacao="remido")
        aberta = self.receber(self.irmaos[0], vencimento=date(2026, 11, 10), competencia=date(2026, 11, 1))
        self.entrar(self.tesoureiro)
        url = reverse("mensalidades_serie")
        self.assertContains(self.client.get(url), "Remido (isenção de mensalidade)")
        base = {"plano": remido.pk, "descricao": "Mensalidade", "valor": "", "meses": 2, "primeiro_vencimento": "2026-11-10",
                "irmaos": [self.irmaos[1].pk]}
        r = self.client.post(url, base)  # nenhum Remido marcado
        self.assertContains(r, "marque na lista de Remidos")
        self.client.post(url, {**base, "remidos": [self.irmaos[0].pk]})  # sem valor: não é obrigatório
        isentos = Lancamento.objects.filter(status="isento").order_by("competencia")
        self.assertEqual([(l.irmao, l.descricao, l.valor) for l in isentos],
                         [(self.irmaos[0], "Isento 11/2026", D("0")), (self.irmaos[0], "Isento 12/2026", D("0"))])
        self.assertFalse(Lancamento.objects.filter(irmao=self.irmaos[1]).exists())  # só os Remidos marcados
        aberta.refresh_from_db()
        self.assertEqual(aberta.status, "cancelado")
        self.assertEqual(servicos.saldo_conta(self.caixa), D("0"))
        # isento não aparece como dívida, mas pode ser consultado pelo filtro
        r = self.client.get(reverse("receber_lista"), {"ano": "2026", "situacao": "aberto", "ver": "irmao"})
        self.assertEqual(r.context["resumo"], [])
        r = self.client.get(reverse("receber_lista"), {"ano": "2026", "situacao": "isento", "ver": "lista"})
        self.assertEqual(len(r.context["linhas"]), 2)
        self.assertContains(self.client.get(reverse("relatorio", args=["mensalidades-ano"]), {"ano": 2026}), "Isento")
        # sem o plano Remido, o valor continua obrigatório
        r = self.client.post(url, {**base, "plano": self.plano["1.1"].pk, "remidos": []})
        self.assertContains(r, "Informe o valor de cada mensalidade")

    def test_lojas_existentes_ganham_a_conta_remido(self):
        from importlib import import_module
        from django.apps import apps as registro
        migracao = import_module("gestao.migrations.0008_plano_remido")
        PlanoConta.objects.filter(loja=self.outra, codigo="1.6").delete()
        migracao.criar_plano_remido(registro, None)
        migracao.criar_plano_remido(registro, None)
        self.assertEqual(PlanoConta.objects.filter(loja=self.outra, nome__startswith="Remido").count(), 1)
        self.assertEqual(PlanoConta.objects.filter(loja=self.loja, nome__startswith="Remido").count(), 1)


class RemidoValorZero(Base):
    def test_remido_aparece_com_r_0_00(self):
        Irmao.objects.filter(pk=self.irmaos[0].pk).update(situacao="remido", centavos=D("0.05"))
        Irmao.objects.filter(pk=self.irmaos[1].pk).update(centavos=D("0.06"))
        self.assertEqual(Irmao.objects.get(pk=self.irmaos[0].pk).valor_mensalidade, D("0.00"))
        self.assertEqual(Irmao.objects.get(pk=self.irmaos[1].pk).valor_mensalidade, D("100"))
        self.entrar(self.master)
        ficha = self.client.get(reverse("irmao_ficha", args=[self.irmaos[0].pk]))
        self.assertContains(ficha, "R$ 0,00")
        self.assertContains(ficha, "remido — isento")
        self.assertContains(self.client.get(reverse("minha_area")), "R$ 0,00")  # o master é o Irmão 0
        controle = self.client.get(reverse("relatorio", args=["mensalidades-ano"]), {"ano": self.hoje.year})
        linha = next(l for l in controle.context["resultado"]["linhas"] if l[1] == "Irmão 0")
        self.assertEqual(linha[3], D("0.00"))
        self.assertContains(controle, "R$ 0,00")
        tabela = self.client.get(reverse("relatorio", args=["centavos"])).context["resultado"]["linhas"]
        self.assertEqual(next(l for l in tabela if l[1] == "Irmão 0")[3], D("0.00"))
        self.assertEqual(next(l for l in tabela if l[1] == "Irmão 1")[3], D("100.06"))
        planilha = self.client.get(reverse("relatorio", args=["mensalidades-ano"]), {"ano": self.hoje.year, "formato": "csv"})
        self.assertIn("Irmão 0;Sim;0,00", planilha.content.decode("utf-8-sig"))


class RemidoComValorZero(Base):
    def test_remido_aparece_com_r0_nos_relatorios_e_na_planilha(self):
        import io
        from openpyxl import load_workbook
        Irmao.objects.filter(pk=self.irmaos[0].pk).update(situacao="remido", centavos=D("0.07"))
        self.entrar(self.tesoureiro)
        r = self.client.get(reverse("relatorio", args=["mensalidades-ano"]), {"ano": self.hoje.year})
        linha = next(l for l in r.context["resultado"]["linhas"] if l[1] == "Irmão 0")
        self.assertEqual(linha[3], D("0.00"))
        self.assertContains(r, "R$ 0,00")
        r = self.client.get(reverse("relatorio", args=["centavos"]))
        linha = next(l for l in r.context["resultado"]["linhas"] if l[1] == "Irmão 0")
        self.assertEqual(linha[3], D("0.00"))
        pasta = load_workbook(io.BytesIO(self.client.get(reverse("planilha_financeira")).content))
        linha = next(l for l in pasta["Mensalidade"].iter_rows(values_only=True) if l[1] == "Irmão 0")
        self.assertEqual((linha[4], linha[5]), (0, "REMIDO"))


class LancamentoRemido(Base):
    def test_remido_aceita_zero_ou_valor_proposto(self):
        remido = PlanoConta.objects.get(loja=self.loja, codigo="1.6")
        self.entrar(self.tesoureiro)
        url = reverse("receber_novo")
        pagina = self.client.get(url)
        self.assertContains(pagina, "data-planos-remido")
        self.assertContains(pagina, 'data-proposto="100')
        base = {"descricao": "Isento", "irmao": self.irmaos[0].pk, "plano": remido.pk, "vencimento": "2026-11-10"}
        self.client.post(url, {**base, "valor": "0"})
        isento = Lancamento.objects.get()
        self.assertEqual((isento.status, isento.valor, isento.valor_pago), ("isento", D("0"), D("0")))
        self.client.post(url, {**base, "descricao": "Mensalidade 11/2026", "valor": "100"})  # valor proposto
        aberto = Lancamento.objects.get(valor=D("100"))
        self.assertEqual(aberto.status, "aberto")
        # alterar o isento para um valor faz dele uma conta em aberto; voltar a 0 o torna isento de novo
        self.client.post(reverse("receber_editar", args=[isento.pk]), {**base, "valor": "80"})
        isento.refresh_from_db()
        self.assertEqual((isento.status, isento.valor), ("aberto", D("80")))
        self.client.post(reverse("receber_editar", args=[isento.pk]), {**base, "valor": "0"})
        isento.refresh_from_db()
        self.assertEqual(isento.status, "isento")
        # fora do plano Remido, zero continua proibido
        r = self.client.post(url, {**base, "plano": self.plano["1.1"].pk, "valor": "0"})
        self.assertContains(r, "O valor precisa ser maior que zero")
        self.assertEqual(servicos.saldo_conta(self.caixa), D("0"))


class LojaInicialDoServidor(TestCase):
    def test_cria_a_loja_real_sem_dados_ficticios_e_recebe_a_carga(self):
        from django.core.management import call_command
        call_command("loja_inicial", stdout=StringIO())
        loja = Loja.objects.get()
        self.assertEqual((loja.nome, loja.numero), ("L'AQUILA ROMANA", "3365"))
        self.assertFalse(Irmao.objects.exists())
        self.assertTrue(PlanoConta.objects.filter(loja=loja, codigo="1.6").exists())
        call_command("loja_inicial", stdout=StringIO())  # segunda vez não duplica
        self.assertEqual(Loja.objects.count(), 1)
        call_command("carga_inicial", stdout=StringIO(), stderr=StringIO())
        self.assertEqual(Irmao.objects.filter(loja=loja).count(), 47)


class ArquivosNoBanco(Base):
    """Fotos, brasão e atas ficam no banco de dados (não se perdem quando a hospedagem reinicia)."""

    def _foto(self, lado=2400, nome="foto.jpg"):
        from io import BytesIO
        from PIL import Image
        saida = BytesIO()
        Image.new("RGB", (lado, lado // 2), (120, 30, 30)).save(saida, "JPEG", quality=95)
        return SimpleUploadedFile(nome, saida.getvalue(), content_type="image/jpeg")

    def test_foto_fica_no_banco_reduzida_e_abre_so_na_propria_loja(self):
        from io import BytesIO
        from PIL import Image
        from .models import ArquivoGuardado
        irmao = self.irmaos[2]
        irmao.foto = self._foto()
        irmao.save()
        guardado = ArquivoGuardado.objects.get(nome=irmao.foto.name)
        self.assertLessEqual(max(Image.open(BytesIO(bytes(guardado.conteudo))).size), 1600)
        self.entrar(self.master)
        resposta = self.client.get(reverse("arquivo", args=[irmao.foto.name]))
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta["Content-Type"], "image/jpeg")
        self.assertIn("max-age", resposta["Cache-Control"])
        self.entrar(self.master_b)
        self.assertEqual(self.client.get(reverse("arquivo", args=[irmao.foto.name])).status_code, 404)

    def test_trocar_e_excluir_limpa_o_arquivo_antigo(self):
        from .models import ArquivoGuardado
        irmao = Irmao.objects.create(loja=self.loja, nome="Com foto")
        irmao.foto = self._foto(lado=300)
        irmao.save()
        primeira = irmao.foto.name
        irmao.foto = self._foto(lado=320, nome="nova.jpg")
        irmao.save()
        self.assertFalse(ArquivoGuardado.objects.filter(nome=primeira).exists())
        self.assertTrue(ArquivoGuardado.objects.filter(nome=irmao.foto.name).exists())
        segunda = irmao.foto.name
        irmao.delete()
        self.assertFalse(ArquivoGuardado.objects.filter(nome=segunda).exists())

    def test_arquivo_antigo_do_disco_vai_para_o_banco(self):
        from pathlib import Path
        from django.core.management import call_command
        from .models import ArquivoGuardado
        with tempfile.TemporaryDirectory() as pasta, override_settings(MEDIA_ROOT=pasta):
            alvo = Path(pasta) / f"loja_{self.loja.pk}" / "atas" / "velha.pdf"
            alvo.parent.mkdir(parents=True)
            alvo.write_bytes(b"%PDF-1.4 teste")
            self.entrar(self.secretario)
            nome = f"loja_{self.loja.pk}/atas/velha.pdf"
            self.assertEqual(self.client.get(reverse("arquivo", args=[nome])).status_code, 200)
            call_command("arquivos_para_banco", stdout=StringIO())
            self.assertEqual(bytes(ArquivoGuardado.objects.get(nome=nome).conteudo), b"%PDF-1.4 teste")


class Monitor(TestCase):
    def test_saude_responde_sem_login_e_sem_banco(self):
        with self.assertNumQueries(0):
            resposta = self.client.get("/saude/")
        self.assertEqual(resposta.content, b"ok")


class EmailPelaInternet(TestCase):
    """No Render gratuito o SMTP é bloqueado: o e-mail sai pela API (HTTPS) do Resend ou do Brevo."""

    def _enviar(self, backend):
        from django.core.mail import EmailMultiAlternatives
        enviados = []

        class Resposta:
            status = 200
            def __enter__(self): return self
            def __exit__(self, *a): return False

        def falso_urlopen(pedido, timeout=0):
            enviados.append((pedido.full_url, dict(pedido.header_items()), json.loads(pedido.data)))
            return Resposta()

        with mock.patch("gestao.email_http.urllib.request.urlopen", falso_urlopen):
            m = EmailMultiAlternatives("Assunto", "Texto", "Loja <nao-responda@ordofraternus.com.br>",
                                       ["irmao@exemplo.com.br"], connection=backend)
            m.attach_alternative("<p>Texto</p>", "text/html")
            self.assertEqual(backend.send_messages([m]), 1)
        return enviados[0]

    @override_settings(RESEND_API_KEY="re_teste")
    def test_resend(self):
        from .email_http import ResendBackend
        url, cab, corpo = self._enviar(ResendBackend())
        self.assertEqual(url, "https://api.resend.com/emails")
        self.assertEqual(cab["Authorization"], "Bearer re_teste")
        self.assertEqual(corpo["to"], ["irmao@exemplo.com.br"])
        self.assertEqual(corpo["html"], "<p>Texto</p>")

    @override_settings(BREVO_API_KEY="xkeysib-teste")
    def test_brevo(self):
        from .email_http import BrevoBackend
        url, cab, corpo = self._enviar(BrevoBackend())
        self.assertEqual(url, "https://api.brevo.com/v3/smtp/email")
        self.assertEqual(cab["Api-key"], "xkeysib-teste")
        self.assertEqual(corpo["sender"], {"email": "nao-responda@ordofraternus.com.br", "name": "Loja"})
